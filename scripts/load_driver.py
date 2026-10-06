"""Synthetic authenticated workload executed in its own disposable driver container."""

import hashlib
import io
import json
import os
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from threading import Lock

import httpx
import psycopg
from docx import Document
from load_fixture import MARKER
from load_metrics import distribution
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from sqlalchemy.engine import make_url

PROFILE = {"id": "synthetic-mixed-office-five-jobs-v1", "workspaces": 5,
           "formats": ["txt", "docx", "xlsx", "pdf"], "paragraphs_per_document": 64,
           "fresh_waves": 3, "concurrency": 5, "admission_capacity": 10,
           "provider": "controlled exact-quote fixture; no inference", "gate_timeout_seconds": 60,
           "http_timeout_seconds": 30, "terminal_wait_seconds": 60, "induced_lock_hold_seconds": 0.25}


def check(condition, code):
    if not condition:
        raise ValueError(code)


def wait_for(predicate, seconds=30):
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        result = predicate()
        if result:
            return result
        time.sleep(.02)
    raise TimeoutError("Synthetic state did not arrive within the declared limit")


def documents(workspace):
    lines = [f"SYNTHETIC FIXTURE {workspace} paragraph {index}: payment odeme evidence for load testing only. No legal authority."
             for index in range(PROFILE["paragraphs_per_document"])]
    yield "txt", "\n\n".join(lines).encode()
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    out = io.BytesIO()
    doc.save(out)
    yield "docx", out.getvalue()
    sheet = Workbook()
    for line in lines:
        sheet.active.append([line, "synthetic", 100])
    out = io.BytesIO()
    sheet.save(out)
    sheet.close()
    yield "xlsx", out.getvalue()
    out = io.BytesIO()
    pdf = canvas.Canvas(out, invariant=1)
    for start in range(0, len(lines), 32):
        text = pdf.beginText(30, 800)
        text.setFont("Helvetica", 8)
        for line in lines[start:start + 32]:
            text.textLine(line)
        pdf.drawText(text)
        pdf.showPage()
    pdf.save()
    yield "pdf", out.getvalue()


class Workload:
    def __init__(self, client, database, report):
        self.client, self.database, self.report = client, database, report
        self.metrics, self.metrics_lock = {}, Lock()
        self.stage = "seed"

    def call(self, method, path, *, metric=None, expected=200, **kwargs):
        start = time.monotonic()
        response = self.client.request(method, path, **kwargs)
        elapsed = time.monotonic() - start
        if metric:
            with self.metrics_lock:
                self.metrics.setdefault(metric, []).append(elapsed)
        check(response.status_code == expected, "Unexpected HTTP status")
        return response.json()

    def state(self):
        return self.call("GET", "/_load/state")

    def gate(self, held):
        return self.call("POST", "/_load/gate/" + ("hold" if held else "release"))

    def submit(self, matter):
        return self.call("POST", f"/api/v1/matters/{matter}/research", metric="admission_seconds", expected=202,
                         json={"question": "SYNTHETIC payment odeme evidence"})["id"]

    def cancel(self, matter, ident, metric):
        return self.call("POST", f"/api/v1/matters/{matter}/research/{ident}/cancel", metric=metric)

    def terminal(self, matter, ident):
        def finished():
            row = self.call("GET", f"/api/v1/matters/{matter}/research/{ident}")
            return row if row["status"] not in {"queued", "running", "cancelling"} else None
        return wait_for(finished, PROFILE["terminal_wait_seconds"])

    def upload(self, matter, suffix, content, *, metric):
        value = self.call("POST", f"/api/v1/matters/{matter}/documents", metric=metric, expected=201,
                          files={"file": (f"SYNTHETIC-load.{suffix}", content, "application/octet-stream")})
        check(value["status"] == "needs_review" and value["passage_count"] > 0, "Document extraction failed")
        check(value["sha256"] == hashlib.sha256(content).hexdigest(), "Uploaded bytes differ from the measured corpus")
        return value["id"]

    def counts(self):
        with psycopg.connect(**self.database) as connection:
            return connection.execute("SELECT (SELECT count(*) FROM records), (SELECT count(*) FROM audit)").fetchone()

    def run(self):
        workspaces, corpus = [], []
        seed_started = time.monotonic()
        for index in range(5):
            matter = self.call("POST", "/api/v1/matters", expected=201, json={"title": f"SYNTHETIC workspace {index}", "domain": "contracts"})["id"]
            workspaces.append(matter)
            for suffix, content in documents(index):
                self.upload(matter, suffix, content, metric="seed_upload_seconds")
                corpus.append({"format": suffix, "bytes": len(content), "sha256": hashlib.sha256(content).hexdigest()})
        self.report["corpus"] = {"files": corpus, "seed_seconds": round(time.monotonic() - seed_started, 6)}
        self.stage = "capacity"
        self.gate(True)
        with ThreadPoolExecutor(5) as pool:
            active = list(pool.map(self.submit, workspaces))
            wait_for(lambda: self.state()["provider"]["active"] == 5)
            queued = list(pool.map(self.submit, workspaces))
            current = self.state()
            check(current["jobs"] == {"admitted": 10, "queued": 5, "running": 5}, "Incorrect admission bound")
            before = self.counts()
            self.call("POST", f"/api/v1/matters/{workspaces[0]}/research", metric="rejection_seconds", expected=429,
                      json={"question": "SYNTHETIC overflow"})
            check(self.counts() == before, "Rejected admission persisted a record or audit")
            cancellations = list(pool.map(lambda pair: self.cancel(*pair, "queued_cancel_seconds"), zip(workspaces, queued, strict=True)))
            check(all(row["status"] == "cancelled" for row in cancellations), "Queued cancellation not acknowledged")
            current = self.state()
            check(current["provider"]["calls"] == 5 and current["jobs"]["admitted"] == 5, "Cancelled queued jobs consumed execution")
            self.report["checks"]["five_workers_ten_admissions_and_clean_overflow"] = True
            self.report["checks"]["queued_cancellations_removed_without_execution"] = True

            self.stage = "postgres_lock_wait"
            with psycopg.connect(**self.database) as blocker, psycopg.connect(**self.database, autocommit=True) as observer:
                blocker.execute("SELECT id FROM records WHERE id = %s FOR UPDATE", (active[0],))
                pid = blocker.info.backend_pid
                started = time.monotonic()
                cancel = pool.submit(self.cancel, workspaces[0], active[0], "running_cancel_with_induced_lock_seconds")
                wait_for(lambda: observer.execute(
                    "SELECT EXISTS (SELECT 1 FROM pg_stat_activity WHERE datname = current_database() "
                    "AND %s = ANY(pg_blocking_pids(pid)))", (pid,)).fetchone()[0], seconds=5)
                observed = time.monotonic()
                time.sleep(PROFILE["induced_lock_hold_seconds"])
                blocker.commit()
                released = time.monotonic()
                cancelled = cancel.result(timeout=10)
                check(cancelled["status"] == "cancelling", "Running stop acknowledged before provider exit")
                self.report["lock_wait"] = {"real_postgres_blocker_observed": True,
                    "request_to_observed_wait_seconds": round(observed - started, 6),
                    "observed_wait_to_release_seconds": round(released - observed, 6),
                    "release_to_cancel_response_seconds": round(time.monotonic() - released, 6)}
            check(self.state()["jobs"]["admitted"] == 5, "Running cancellation released capacity early")
            self.report["checks"]["running_cancellation_retains_capacity_until_unwind"] = True

            self.stage = "overlapping_ingestion"
            check(self.state()["provider"]["active"] == 5, "Research no longer overlaps ingestion")
            additions = list(pool.map(lambda matter: self.upload(matter, "txt", b"SYNTHETIC added payment evidence during research.\n" * 100,
                                                                metric="overlapping_upload_seconds"), workspaces))
            check(self.state()["provider"]["active"] == 5, "Provider exited before ingestion completed")
            self.report["checks"]["five_uploads_complete_while_five_jobs_active"] = True
            self.gate(False)
            outcomes = list(pool.map(lambda pair: self.terminal(*pair), zip(workspaces, active, strict=True)))
            check(outcomes[0]["status"] == "cancelled" and "product_id" not in outcomes[0], "Cancelled job published output")
            check(all(row["status"] == "completed" for row in outcomes[1:]), "Research did not complete after ingestion")
            for matter, row in zip(workspaces[1:], outcomes[1:], strict=True):
                product = self.call("GET", f"/api/v1/matters/{matter}/products/{row['product_id']}")
                check(product["status"] == "stale", "Changed evidence did not make in-flight work stale")
            wait_for(lambda: self.state()["jobs"]["admitted"] == 0)
            self.report["checks"]["cancellation_blocks_output_and_changed_evidence_marks_outputs_stale"] = True

            self.stage = "fresh_waves"
            elapsed, queue_times, executions = [], [], []
            waves = []
            for _ in range(PROFILE["fresh_waves"]):
                start = time.monotonic()
                ids = list(pool.map(self.submit, workspaces))
                rows = list(pool.map(lambda pair: self.terminal(*pair), zip(workspaces, ids, strict=True)))
                waves.append(time.monotonic() - start)
                for matter, addition, row in zip(workspaces, additions, rows, strict=True):
                    check(row["status"] == "completed", "Fresh job did not complete")
                    product = self.call("GET", f"/api/v1/matters/{matter}/products/{row['product_id']}")
                    check(product["status"] == "needs_review" and addition in product["snapshots"]["sources"], "Fresh work missed current evidence")
                    evidence = {item["id"]: item["text"] for item in product["evidence"]}
                    check(bool(product["claims"]) and all(claim["text"] in evidence[ref] for claim in product["claims"] for ref in claim["evidence_ids"]), "Quote evidence mismatch")
                    created, began, finished = (datetime.fromisoformat(row[field]) for field in ("created_at", "started_at", "finished_at"))
                    elapsed.append((finished - created).total_seconds())
                    queue_times.append((began - created).total_seconds())
                    executions.append((finished - began).total_seconds())
                wait_for(lambda: self.state()["jobs"]["admitted"] == 0)
            self.report["fresh_jobs"] = {"count": len(elapsed), "job_seconds": distribution(elapsed),
                "queue_seconds": distribution(queue_times), "execution_seconds": distribution(executions),
                "wave_seconds": waves, "jobs_per_second": round(len(elapsed) / sum(waves), 6)}
            self.report["checks"]["fresh_waves_use_current_evidence_and_exact_quotes"] = True
        final = self.state()
        check(final["provider"]["peak"] == 5 and final["provider"]["active"] == 0 and final["jobs"]["admitted"] == 0, "Workers did not drain")
        self.report["final_counters"] = final
        self.report["checks"]["worker_concurrency_never_exceeds_five"] = True


def main():
    report = {"status": "failed", "profile": PROFILE, "checks": {}}
    workload = None
    try:
        url = make_url(os.environ["LA_DATABASE_URL"])
        check(os.environ.get("LA_LOAD_DRILL") == MARKER and url.host == "postgres"
              and url.database == "lawyer_load_drill" and not url.query, "Wrong synthetic target")
        database = {"host": url.host, "port": url.port, "user": url.username, "password": url.password, "dbname": url.database,
                    "connect_timeout": 5, "options": "-c statement_timeout=15000 -c lock_timeout=10000"}
        with httpx.Client(base_url="http://api:8000", timeout=30, trust_env=False,
                          headers={"X-Load-Token": os.environ["LA_LOAD_CONTROL_TOKEN"]}) as client:
            workload = Workload(client, database, report)
            login = workload.call("POST", "/api/v1/auth/login", json={"username": "demo", "password": "demo-local-only"})
            client.headers["X-CSRF-Token"] = login["csrf_token"]
            try:
                workload.run()
                report["status"] = "passed"
            finally:
                workload.gate(False)
    except Exception as error:
        report.update(status="failed", failure_stage=workload.stage if workload else "configuration", failure_type=type(error).__name__)
    if workload:
        report["request_seconds"] = {name: distribution(values) for name, values in workload.metrics.items()}
    print(json.dumps(report, sort_keys=True))
    return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
