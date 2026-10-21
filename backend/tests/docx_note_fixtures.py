"""Invented OOXML note parts; never opens user files or external resources."""

from copy import deepcopy

from docx import Document
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree


def note_reference(paragraph, kind, identifier):
    ref = OxmlElement("w:" + kind + "Reference")
    if identifier is not None:
        ref.set(qn("w:id"), str(identifier))
    paragraph.add_run()._r.append(ref)
    return ref


def note_part(document, kind, entries):
    """Entries are (package ID, type or None, text or a synthetic block document)."""
    root = OxmlElement("w:" + kind + "s")
    for identifier, note_type, content in entries:
        node = OxmlElement("w:" + kind)
        node.set(qn("w:id"), str(identifier))
        if note_type is not None:
            node.set(qn("w:type"), note_type)
        if isinstance(content, str):
            block_doc = Document()
            p = block_doc.add_paragraph(content)
            p.add_run()._r.insert(0, OxmlElement("w:" + kind + "Ref"))
        else:
            block_doc = content
        for block in block_doc.element.body:
            if block.tag != qn("w:sectPr"):
                node.append(deepcopy(block))
        root.append(node)
    part = Part(PackURI("/word/" + kind + "s.xml"),
                getattr(CT, "WML_" + kind.upper() + "S"), etree.tostring(root), document.part.package)
    document.part.relate_to(part, getattr(RT, kind.upper() + "S"))
    return part
