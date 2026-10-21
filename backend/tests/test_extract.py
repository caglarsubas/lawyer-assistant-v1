import zipfile
from email.message import EmailMessage

import pytest
from pptx import Presentation

from app.extract import extract


def write(tmp_path, suffix, content):
    path = tmp_path / ("input" + suffix)
    path.write_bytes(content)
    return path


def test_html_never_returns_script_or_embedded_object(tmp_path):
    path = write(
        tmp_path,
        ".html",
        b"<p>Contract</p><script>steal()</script><style>hide</style><iframe>remote</iframe><p>Term</p>",
    )
    output = extract(path, ".html")
    assert output["passages"][0]["text"] == "Contract\nTerm"
    assert output["warnings"]


def test_email_attachment_is_explicitly_unprocessed(tmp_path):
    mail = EmailMessage()
    mail["Subject"] = "İş sözleşmesi"
    mail["From"] = "example@example.invalid"
    mail.set_content("Ödeme 30 gün içinde yapılır.")
    mail.add_attachment(
        b"Never parse this automatically",
        maintype="application",
        subtype="octet-stream",
        filename="private.bin",
    )
    result = extract(write(tmp_path, ".eml", mail.as_bytes()), ".eml")
    assert any("30 gün" in p["text"] for p in result["passages"])
    assert not any("Never parse" in p["text"] for p in result["passages"])
    assert any("1 ek işlenmedi" in w for w in result["warnings"])


def test_presentation_has_slide_locators_and_omission_warning(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Ticari uyuşmazlık"
    slide.placeholders[1].text = "İddia ile bulgu ayrılır."
    path = tmp_path / "input.pptx"
    deck.save(path)
    result = extract(path, ".pptx")
    assert result["page_count"] == 1
    assert result["passages"][0]["locator"].startswith("Slayt 1")
    assert any("İddia" in p["text"] for p in result["passages"])
    assert result["warnings"]


@pytest.mark.parametrize("suffix", [".odt", ".ods", ".odp"])
def test_opendocument_bounded_paragraphs_and_repeated_cells(tmp_path, suffix):
    path = tmp_path / ("input" + suffix)
    content = """<office xmlns:t="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
        xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0">
        <table:table-row table:number-rows-repeated="999999999"><t:p>İşlem</t:p></table:table-row></office>"""
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("content.xml", content)
    result = extract(path, suffix)
    assert [p["text"] for p in result["passages"]] == ["İşlem"]
    assert any("çoğaltılmadı" in w for w in result["warnings"])


def test_xml_entities_and_office_macros_rejected(tmp_path):
    path = tmp_path / "input.odt"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(
            "content.xml", '<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///etc/passwd">]><x>&secret;</x>'
        )
    from defusedxml.common import DefusedXmlException

    with pytest.raises(DefusedXmlException):
        extract(path, ".odt")
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/vbaProject.bin", b"macro")
    with pytest.raises(ValueError, match="Macros"):
        extract(path, ".docx")


def test_rtf_unicode_and_invalid_header(tmp_path):
    result = extract(write(tmp_path, ".rtf", rb"{\rtf1\ansi \u304?\u351? hukuku}"), ".rtf")
    assert result["passages"][0]["text"] == "İş hukuku"
    with pytest.raises(ValueError, match="RTF"):
        extract(write(tmp_path, ".rtf", b"not rtf"), ".rtf")


def test_legacy_failure_and_supplied_udf_remain_visible(tmp_path):
    for suffix in (".doc", ".xls", ".msg", ".udf"):
        with pytest.raises(Exception):
            extract(write(tmp_path, suffix, b"not a valid document"), suffix)


def test_xlsx_formula_is_literal_and_csv_keeps_row_locator(tmp_path):
    from openpyxl import Workbook

    book = Workbook()
    book.active["A1"] = '=WEBSERVICE("https://example.invalid")'
    path = tmp_path / "input.xlsx"
    book.save(path)
    result = extract(path, ".xlsx")
    assert "=WEBSERVICE" in result["passages"][0]["text"]
    assert any("çalıştırılmadı" in w for w in result["warnings"])
    result = extract(write(tmp_path, ".csv", "isim,tutar\nörnek,30".encode()), ".csv")
    assert result["passages"][1] == {"locator": "Satır 2", "text": "örnek | 30"}


def test_docx_interleaves_tables_with_clauses_without_renumbering(tmp_path):
    from docx import Document

    from app.extract import DOCX_BODY_PROFILE

    doc = Document()
    doc.add_paragraph("SENTETİK: Tablodaki tutar kabul edilmedi.")
    doc.add_table(rows=1, cols=2).rows[0].cells[0].text = "1.234,56 TL"
    doc.tables[0].cell(0, 1).text = "14.06.2020"
    doc.add_paragraph("")  # Empty body paragraph still has its own locator.
    doc.add_paragraph("SENTETİK: Önceki tablodaki tarih kesinleşmedi.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "İkinci tablo: TEST-0042"
    doc.add_paragraph("SENTETİK: Son hüküm.")
    path = tmp_path / "input.docx"
    doc.save(path)
    original = path.read_bytes()
    result = extract(path, ".docx")
    assert result["passages"] == [
        {"locator": "Paragraf 1", "text": "SENTETİK: Tablodaki tutar kabul edilmedi."},
        {"locator": "Tablo 1, satır 1", "text": "1.234,56 TL | 14.06.2020"},
        {"locator": "Paragraf 3", "text": "SENTETİK: Önceki tablodaki tarih kesinleşmedi."},
        {"locator": "Tablo 2, satır 1", "text": "İkinci tablo: TEST-0042"},
        {"locator": "Paragraf 4", "text": "SENTETİK: Son hüküm."},
    ]
    assert result["extraction_profile"] == DOCX_BODY_PROFILE
    assert result["passage_count"] == 5 and path.read_bytes() == original


def test_docx_nested_tables_preserve_exact_cell_block_sequence(tmp_path):
    from docx import Document

    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "SENTETİK: Önceki koşul kabul edilmedi."
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "İç tablo: 1.234,56 TL; 14.06.2020."
    cell.add_paragraph("SENTETİK: Ancak TEST-0042 istisnası saklıdır.")
    path = tmp_path / "input.docx"
    doc.save(path)
    original = path.read_bytes()
    result = extract(path, ".docx")
    assert result["passages"] == [
        {"locator": "Tablo 1, satır 1, hücre 1, paragraf 1",
         "text": "SENTETİK: Önceki koşul kabul edilmedi."},
        {"locator": "Tablo 1, satır 1, hücre 1, tablo 1, satır 1",
         "text": "İç tablo: 1.234,56 TL; 14.06.2020."},
        {"locator": "Tablo 1, satır 1, hücre 1, paragraf 3",
         "text": "SENTETİK: Ancak TEST-0042 istisnası saklıdır."},
    ]
    assert path.read_bytes() == original
    assert any("fiziksel XML" in warning for warning in result["warnings"])
    assert any("dipnot" in warning for warning in result["warnings"])


def test_docx_nested_rows_keep_sibling_cells_tables_and_body_order(tmp_path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("ÖNCE")
    row = doc.add_table(rows=2, cols=2).rows[0]
    row.cells[0].text = "HÜCRE-A"
    nested = row.cells[0].add_table(rows=1, cols=1).cell(0, 0)
    nested.text = "İÇ-ÖNCE"
    nested.add_table(rows=1, cols=1).cell(0, 0).text = "DERİN-KOŞUL"
    nested.add_paragraph("İÇ-SONRA")
    row.cells[0].add_table(rows=1, cols=1).cell(0, 0).text = "İKİNCİ-İÇ-TABLO"
    row.cells[1].text = "HÜCRE-B"
    doc.tables[0].cell(1, 0).text = "SATIR-2"
    doc.add_paragraph("SONRA")
    path = tmp_path / "input.docx"
    doc.save(path)
    result = extract(path, ".docx")
    assert [p["text"] for p in result["passages"]] == [
        "ÖNCE", "HÜCRE-A", "İÇ-ÖNCE", "DERİN-KOŞUL", "İÇ-SONRA",
        "İKİNCİ-İÇ-TABLO", "HÜCRE-B", "SATIR-2 |", "SONRA",
    ]
    assert len({p["locator"] for p in result["passages"]}) == result["passage_count"]
    assert "hücre 1, tablo 2" in result["passages"][5]["locator"]
    assert result["passages"][6]["locator"] == "Tablo 1, satır 1, hücre 2, paragraf 1"


@pytest.mark.parametrize("direction", ["horizontal", "vertical"])
def test_docx_merged_cell_content_is_not_duplicated(tmp_path, direction):
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "BİRLEŞMİŞ-ÖZGÜN-KOŞUL"
    table.cell(0, 0).merge(table.cell(0, 1) if direction == "horizontal" else table.cell(1, 0))
    table.cell(1, 1).text = "BAĞIMSIZ-KOŞUL"
    path = tmp_path / "input.docx"
    doc.save(path)
    result = extract(path, ".docx")
    text = "\n".join(p["text"] for p in result["passages"])
    assert text.count("BİRLEŞMİŞ-ÖZGÜN-KOŞUL") == text.count("BAĞIMSIZ-KOŞUL") == 1
    assert any("Birleştirilmiş" in w for w in result["warnings"])


@pytest.mark.parametrize("depth,admitted", [(4, True), (5, False)])
def test_docx_nesting_budget_never_returns_partial_content(tmp_path, depth, admitted):
    from docx import Document

    doc = Document()
    doc.add_paragraph("BAŞLANGIÇ; kısmi başarı olarak dönmemeli.")
    parent = doc
    for _ in range(depth):
        parent = parent.add_table(rows=1, cols=1).cell(0, 0)
    parent.text = "EN-DERİN"
    path = tmp_path / "input.docx"
    doc.save(path)
    if admitted:
        result = extract(path, ".docx")
        assert result["passages"][-1]["text"] == "EN-DERİN"
        assert max(len(p["locator"]) for p in result["passages"]) <= 400
    else:
        with pytest.raises(ValueError, match="^DOCX table nesting limit exceeded$"):
            extract(path, ".docx")


def test_docx_structural_budget_counts_empty_cells(tmp_path, monkeypatch):
    from docx import Document

    monkeypatch.setattr("app.extract.MAX_DOCX_BLOCKS", 5)
    doc = Document()
    doc.add_table(rows=1, cols=4)  # Six visited blocks, even with no text.
    path = tmp_path / "input.docx"
    doc.save(path)
    with pytest.raises(ValueError, match="^DOCX structural work limit exceeded$"):
        extract(path, ".docx")


def test_docx_structural_budget_counts_empty_paragraphs_in_plain_cells(tmp_path, monkeypatch):
    from docx import Document

    monkeypatch.setattr("app.extract.MAX_DOCX_BLOCKS", 6)
    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    for _ in range(4):
        cell.add_paragraph("")
    path = tmp_path / "input.docx"
    doc.save(path)
    with pytest.raises(ValueError, match="^DOCX structural work limit exceeded$"):
        extract(path, ".docx")


def test_docx_grid_span_is_not_expanded_into_repeated_cells(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    cell = doc.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "SENTETİK: Tek fiziksel koşul."
    span = OxmlElement("w:gridSpan")
    span.set(qn("w:val"), "1000000000")
    cell._tc.get_or_add_tcPr().append(span)
    path = tmp_path / "input.docx"
    doc.save(path)
    assert extract(path, ".docx")["passages"] == [
        {"locator": "Tablo 1, satır 1", "text": "SENTETİK: Tek fiziksel koşul."},
    ]


def test_docx_header_footer_definitions_keep_section_origin_and_block_order(tmp_path):
    from docx import Document
    from docx.enum.section import WD_SECTION_START

    doc = Document()
    doc.add_paragraph("SENTETİK: Gövde; bu borç kabul edilmedi.")
    header = doc.sections[0].header
    header.paragraphs[0].text = "SENTETİK: Üst koşul; 14.06.2020."
    cell = header.add_table(rows=1, cols=1, width=1000).cell(0, 0)
    cell.text = "SENTETİK: Dış koşul."
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "SENTETİK: 1.234,56 TL istisnası."
    cell.add_paragraph("SENTETİK: İç tablodan sonra.")
    header.add_paragraph("SENTETİK: Üst son.")
    doc.sections[0].footer.paragraphs[0].text = "SENTETİK: Alt koşul; TEST-0042."
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    path = tmp_path / "input.docx"
    doc.save(path)
    original = path.read_bytes()
    result = extract(path, ".docx")
    header_passages = [p for p in result["passages"] if "üstbilgi" in p["locator"]]
    assert [p["text"] for p in header_passages] == [
        "SENTETİK: Üst koşul; 14.06.2020.", "SENTETİK: Dış koşul.",
        "SENTETİK: 1.234,56 TL istisnası.", "SENTETİK: İç tablodan sonra.",
        "SENTETİK: Üst son.",
    ] * 2
    assert header_passages[0]["locator"] == (
        "Bölüm 1, üstbilgi tanımı (varsayılan; kaynak bölüm 1), paragraf 1"
    )
    assert header_passages[5]["locator"].startswith(
        "Bölüm 2, üstbilgi tanımı (varsayılan; kaynak bölüm 1)"
    )
    assert [p["text"] for p in result["passages"] if "altbilgi" in p["locator"]] == [
        "SENTETİK: Alt koşul; TEST-0042.", "SENTETİK: Alt koşul; TEST-0042.",
    ]
    assert result["passages"][0]["text"].startswith("SENTETİK: Gövde;")
    assert path.read_bytes() == original and result["page_count"] is None
    assert any("sayfa" in w and "tanımları" in w for w in result["warnings"])


def test_docx_empty_override_stops_header_inheritance_without_replacing_footer(tmp_path):
    from docx import Document
    from docx.enum.section import WD_SECTION_START

    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "SENTETİK: Eski üst koşul."
    doc.sections[0].footer.paragraphs[0].text = "SENTETİK: Ortak alt koşul."
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    blank = doc.add_section(WD_SECTION_START.NEW_PAGE)
    blank.header.is_linked_to_previous = False
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    path = tmp_path / "input.docx"
    doc.save(path)
    passages = extract(path, ".docx")["passages"]
    assert len([p for p in passages if "üstbilgi" in p["locator"]]) == 2
    assert len([p for p in passages if "altbilgi" in p["locator"]]) == 4


def test_docx_first_even_and_default_definitions_do_not_claim_page_activation(tmp_path):
    from docx import Document

    doc = Document()
    # Definitions can exist while the first/even-page rendering switches are off.
    for name in ["header", "first_page_header", "even_page_header",
                 "footer", "first_page_footer", "even_page_footer"]:
        getattr(doc.sections[0], name).paragraphs[0].text = "SENTETİK: " + name
    path = tmp_path / "input.docx"
    doc.save(path)
    result = extract(path, ".docx")
    assert len(result["passages"]) == 6
    assert {p["text"] for p in result["passages"]} == {
        "SENTETİK: " + name for name in ["header", "first_page_header", "even_page_header",
                                      "footer", "first_page_footer", "even_page_footer"]
    }
    assert all("tanımı" in p["locator"] for p in result["passages"])
    assert any("ilk sayfa" in p["locator"] for p in result["passages"])
    assert any("çift sayfa" in p["locator"] for p in result["passages"])
    assert any("etkinliği" in w for w in result["warnings"])


@pytest.mark.parametrize("broken", [
    "external", "wrong_kind", "duplicate", "unknown_variant", "missing", "wrong_root",
])
def test_docx_invalid_story_reference_rejects_complete_extraction(tmp_path, broken):
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    doc.add_paragraph("SENTETİK: Bu kısmi gövde döndürülmemeli.")
    section = doc.sections[0]
    section.header.paragraphs[0].text = "SENTETİK: Tanım."
    ref = section._sectPr.headerReference_lst[0]
    if broken == "external":
        ref.set(qn("r:id"), doc.part.relate_to("https://example.invalid/no-fetch", RT.HEADER, True))
    elif broken == "wrong_kind":
        section.footer.paragraphs[0].text = "SENTETİK: Yanlış tür."
        ref.set(qn("r:id"), section._sectPr.footerReference_lst[0].get(qn("r:id")))
    elif broken == "duplicate":
        copy = OxmlElement("w:headerReference")
        copy.attrib.update(ref.attrib)
        section._sectPr.append(copy)
    elif broken == "unknown_variant":
        ref.set(qn("w:type"), "unsupported")
    elif broken == "missing":
        ref.set(qn("r:id"), "rIdMissing")
    else:
        section.header._element.tag = qn("w:ftr")
    path = tmp_path / "input.docx"
    doc.save(path)
    with pytest.raises(ValueError, match="^DOCX header/footer reference is invalid$"):
        extract(path, ".docx")


@pytest.mark.parametrize("limit", ["depth", "work"])
def test_docx_header_tables_share_body_depth_and_structural_limits(tmp_path, monkeypatch, limit):
    from docx import Document

    doc = Document()
    doc.add_paragraph("SENTETİK: Başlangıç.")
    header = doc.sections[0].header
    if limit == "depth":
        parent = header.add_table(rows=1, cols=1, width=1000).cell(0, 0)
        for _ in range(4):
            parent = parent.add_table(rows=1, cols=1).cell(0, 0)
        message = "DOCX table nesting limit exceeded"
    else:
        monkeypatch.setattr("app.extract.MAX_DOCX_BLOCKS", 8)
        for _ in range(10):
            header.add_paragraph("")
        message = "DOCX structural work limit exceeded"
    path = tmp_path / "input.docx"
    doc.save(path)
    with pytest.raises(ValueError, match="^" + message + "$"):
        extract(path, ".docx")


def test_docx_body_and_stories_share_total_text_limit(tmp_path, monkeypatch):
    from docx import Document

    monkeypatch.setattr("app.extract.MAX_TEXT", 50)
    doc = Document()
    doc.add_paragraph("SENTETİK: Gövde 1234567890.")
    doc.sections[0].header.paragraphs[0].text = "SENTETİK: Üst koşul 1234567890."
    path = tmp_path / "input.docx"
    doc.save(path)
    with pytest.raises(ValueError, match="^Extracted text limit exceeded$"):
        extract(path, ".docx")


@pytest.mark.parametrize("tag", ["ins", "del", "moveFrom", "moveTo"])
def test_docx_story_revisions_are_reported_without_claiming_an_accepted_version(tmp_path, tag):
    from docx import Document
    from docx.oxml import OxmlElement

    doc = Document()
    doc.add_paragraph("SENTETİK: Gövde.")
    doc.sections[0].header.paragraphs[0]._p.append(OxmlElement("w:" + tag))
    path = tmp_path / "input.docx"
    doc.save(path)
    assert "İzlenen değişiklikler var; kabul edilmiş sürüm doğrulanmadı." in extract(
        path, ".docx"
    )["warnings"]


def test_docx_field_instruction_is_not_misclassified_as_a_revision(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement

    doc = Document()
    doc.add_paragraph("SENTETİK: Gövde.")
    field = OxmlElement("w:instrText")
    field.text = "PAGE"
    doc.sections[0].footer.paragraphs[0].add_run()._r.append(field)
    path = tmp_path / "input.docx"
    doc.save(path)
    warnings = extract(path, ".docx")["warnings"]
    assert not any("İzlenen değişiklikler var" in w for w in warnings)
    assert any("alan değerleri doğrulanmadı" in w for w in warnings)


def test_docx_notes_keep_exact_reference_context_kind_and_content_order(tmp_path):
    from docx import Document
    from docx_note_fixtures import note_part, note_reference

    doc = Document()
    first = doc.add_paragraph("SENTETİK: Borç kabul edilmedi.")
    note_reference(first, "footnote", 7)
    second = doc.add_paragraph("SENTETİK: Ancak ayrı istisna saklıdır.")
    note_reference(second, "endnote", 7)
    third = doc.add_paragraph("SENTETİK: İlk koşula tekrar başvuru.")
    note_reference(third, "footnote", 7)
    blocks = Document()
    blocks.add_paragraph("SENTETİK: 14.06.2020 koşulu.")
    cell = blocks.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "SENTETİK: Dış koşul."
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "SENTETİK: 1.234,56 TL kabul edilmedi."
    cell.add_paragraph("SENTETİK: İç istisna saklıdır.")
    blocks.add_paragraph("SENTETİK: Dipnot sonu; TEST-0042.")
    note_part(doc, "footnote", [(7, None, blocks)])
    note_part(doc, "endnote", [(7, "normal", "SENTETİK: Sonnot ayrı koşuldur.")])
    path = tmp_path / "input.docx"
    doc.save(path)
    original = path.read_bytes()
    result = extract(path, ".docx")
    footnotes = [p for p in result["passages"] if p["locator"].startswith("Dipnot kimlik 7")]
    assert [p["text"] for p in footnotes] == [
        "SENTETİK: 14.06.2020 koşulu.", "SENTETİK: Dış koşul.",
        "SENTETİK: 1.234,56 TL kabul edilmedi.", "SENTETİK: İç istisna saklıdır.",
        "SENTETİK: Dipnot sonu; TEST-0042.",
    ] * 2
    assert footnotes[0]["locator"] == "Dipnot kimlik 7, başvuru 1 (Paragraf 1), paragraf 1"
    assert footnotes[5]["locator"].startswith("Dipnot kimlik 7, başvuru 3 (Paragraf 3)")
    assert result["passages"][-1] == {
        "locator": "Sonnot kimlik 7, başvuru 2 (Paragraf 2), paragraf 1",
        "text": "SENTETİK: Sonnot ayrı koşuldur.",
    }
    assert result["passages"][0] == {"locator": "Paragraf 1", "text": first.text}
    assert path.read_bytes() == original and result["page_count"] is None
    assert any("paket kimliği" in w and "gösterilen" in w for w in result["warnings"])
