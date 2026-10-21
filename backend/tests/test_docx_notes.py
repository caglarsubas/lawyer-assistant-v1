"""Synthetic note identity, source-context and resource-bound regressions."""

import zipfile
from copy import deepcopy

import pytest
from defusedxml.common import DefusedXmlException
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx_note_fixtures import note_part, note_reference
from lxml import etree

from app.extract import extract


def save_and_extract(doc, tmp_path):
    path = tmp_path / "invented.docx"
    doc.save(path)
    before = path.read_bytes()
    result = extract(path, ".docx")
    assert path.read_bytes() == before
    return result


@pytest.mark.parametrize("kind,label", [("footnote", "Dipnot"), ("endnote", "Sonnot")])
@pytest.mark.parametrize("identifier", [0, -5, 2147483647])
def test_note_identity_is_not_display_number_or_separator_identity(tmp_path, kind, label, identifier):
    doc = Document()
    p = doc.add_paragraph("SENTETİK: Koşul ")
    ref = note_reference(p, kind, identifier)
    ref.set(qn("w:customMarkFollows"), "1")
    p.add_run("†")
    note_part(doc, kind, [(identifier, "normal", "SENTETİK: Not koşulu.")])
    result = save_and_extract(doc, tmp_path)
    assert result["passages"] == [
        {"locator": "Paragraf 1", "text": "SENTETİK: Koşul †"},
        {"locator": f"{label} kimlik {identifier}, başvuru 1 (Paragraf 1), paragraf 1",
         "text": "SENTETİK: Not koşulu."},
    ]


def test_separator_types_are_excluded_and_unreferenced_normal_notes_are_explicit(tmp_path):
    doc = Document()
    doc.add_paragraph("SENTETİK: Gövde.")
    note_part(doc, "footnote", [
        (8, "separator", "SYSTEM SEPARATOR: NOT A LEGAL NOTE"),
        (9, "continuationSeparator", "SYSTEM CONTINUATION"),
        (10, "continuationNotice", "SYSTEM NOTICE"),
        (1, None, "SENTETİK: Bağlanmamış not."),
    ])
    result = save_and_extract(doc, tmp_path)
    assert result["passages"][-1] == {
        "locator": "Dipnot kimlik 1, başvuru yok, paragraf 1",
        "text": "SENTETİK: Bağlanmamış not.",
    }
    assert not any("SYSTEM" in p["text"] for p in result["passages"])
    assert any("Başvurusu bulunmayan" in w for w in result["warnings"])


def test_marker_only_paragraph_and_empty_note_are_not_invented_source_text(tmp_path):
    doc = Document()
    note_reference(doc.add_paragraph(""), "footnote", 7)
    note_part(doc, "footnote", [(7, None, "SENTETİK: Yalnız not metni.")])
    result = save_and_extract(doc, tmp_path)
    assert result["passages"] == [{
        "locator": "Dipnot kimlik 7, başvuru 1 (Paragraf 1), paragraf 1",
        "text": "SENTETİK: Yalnız not metni.",
    }]
    doc = Document()
    note_reference(doc.add_paragraph("SENTETİK: Gövde."), "endnote", 2)
    note_part(doc, "endnote", [(2, None, "")])
    result = save_and_extract(doc, tmp_path)
    assert result["passage_count"] == 1
    assert any("Boş not tanımı" in w for w in result["warnings"])


def test_table_and_inherited_story_note_references_retain_their_own_context(tmp_path):
    from docx.enum.section import WD_SECTION_START

    doc = Document()
    p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    p.text = "SENTETİK: Hücre koşulu."
    note_reference(p, "footnote", 7)
    p = doc.sections[0].header.paragraphs[0]
    p.text = "SENTETİK: Tanım koşulu."
    note_reference(p, "endnote", 7)
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    note_part(doc, "footnote", [(7, None, "SENTETİK: Hücre notu.")])
    note_part(doc, "endnote", [(7, None, "SENTETİK: Tanım notu.")])
    result = save_and_extract(doc, tmp_path)
    notes = [p for p in result["passages"] if p["locator"].startswith(("Dipnot", "Sonnot"))]
    assert len(notes) == 3
    assert "başvuru 1 (Tablo 1, satır 1, hücre 1, paragraf 1)" in notes[0]["locator"]
    assert "Bölüm 1, üstbilgi tanımı (varsayılan; kaynak bölüm 1)" in notes[1]["locator"]
    assert "Bölüm 2, üstbilgi tanımı (varsayılan; kaynak bölüm 1)" in notes[2]["locator"]


@pytest.mark.parametrize("identifier", [None, "", "x", "2147483648", "-2147483649", "9" * 100])
def test_malformed_reference_ids_fail_without_partial_success(tmp_path, identifier):
    doc = Document()
    note_reference(doc.add_paragraph("SENTETİK: Döndürülmemeli."), "footnote", identifier)
    note_part(doc, "footnote", [(7, None, "SENTETİK: Not.")])
    with pytest.raises(ValueError, match="^DOCX note definition/reference is invalid$"):
        save_and_extract(doc, tmp_path)


@pytest.mark.parametrize("defect", [
    "missing", "external", "duplicate_part", "wrong_content_type", "wrong_root", "duplicate_id",
    "unknown_type", "separator_reference", "missing_id", "unresolved", "nested_reference", "wrapper",
])
def test_invalid_note_parts_and_links_fail_closed(tmp_path, defect):
    doc = Document()
    note_reference(doc.add_paragraph("SENTETİK: Kısmi gövde dönmemeli."), "footnote", 7)
    entries = [(7, None, "SENTETİK: Not.")]
    if defect == "duplicate_id":
        entries.append(("007", "normal", "SENTETİK: Aynı kimlik."))
    if defect == "unknown_type":
        entries = [(7, "unsupported", "SENTETİK: Not.")]
    if defect == "separator_reference":
        entries = [(7, "separator", "SYSTEM CONTENT MUST NOT BECOME A NOTE")]
    if defect == "unresolved":
        entries = [(8, None, "SENTETİK: Başka not.")]
    part = note_part(doc, "footnote", entries)
    rid = next(r.rId for r in doc.part.rels.values() if r.reltype == RT.FOOTNOTES)
    if defect in {"missing", "external"}:
        del doc.part.rels[rid]
        if defect == "external":
            doc.part.rels.add_relationship(RT.FOOTNOTES, "https://example.invalid/not-fetched", rid, True)
    elif defect == "duplicate_part":
        another = Part(PackURI("/word/alternate.xml"), part.content_type, part.blob, doc.part.package)
        doc.part.relate_to(another, RT.FOOTNOTES)
    elif defect == "wrong_content_type":
        part._content_type = "application/xml"
    elif defect in {"wrong_root", "missing_id", "nested_reference", "wrapper"}:
        root = etree.fromstring(part.blob)
        if defect == "wrong_root":
            root.tag = qn("w:endnotes")
        elif defect == "missing_id":
            del root[0].attrib[qn("w:id")]
        elif defect == "nested_reference":
            ref = OxmlElement("w:endnoteReference")
            ref.set(qn("w:id"), "7")
            root[0][0][0].append(ref)
        else:
            root[0].append(OxmlElement("w:sdt"))
        part._blob = etree.tostring(root)
    with pytest.raises(ValueError, match="^DOCX (note|nested note)"):
        save_and_extract(doc, tmp_path)


def test_duplicate_raw_relationship_ids_are_not_collapsed_into_one_authority(tmp_path):
    doc = Document()
    note_reference(doc.add_paragraph("SENTETİK: Gövde."), "footnote", 7)
    note_part(doc, "footnote", [(7, None, "SENTETİK: Not.")])
    path = tmp_path / "invented.docx"
    doc.save(path)
    with zipfile.ZipFile(path) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    name = "word/_rels/document.xml.rels"
    root = etree.fromstring(members[name])
    root.append(deepcopy(next(r for r in root if r.get("Type") == RT.FOOTNOTES)))
    members[name] = etree.tostring(root)
    with zipfile.ZipFile(path, "w") as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    with pytest.raises(ValueError, match="^DOCX note relationship is invalid$"):
        extract(path, ".docx")


def test_note_xml_dtd_and_external_entities_are_rejected(tmp_path):
    doc = Document()
    doc.add_paragraph("SENTETİK: Gövde.")
    part = note_part(doc, "footnote", [])
    part._blob = (b'<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///etc/passwd">]>'
                  b'<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                  b'&secret;</w:footnotes>')
    with pytest.raises(DefusedXmlException):
        save_and_extract(doc, tmp_path)


@pytest.mark.parametrize("wrapper", ["ins", "sdt"])
def test_omitted_reference_context_is_not_silently_linked_or_made_an_orphan(tmp_path, wrapper):
    doc = Document()
    p = doc.add_paragraph("SENTETİK: Gövde.")
    ref = note_reference(p, "footnote", 7)
    node = OxmlElement("w:" + wrapper)
    if wrapper == "ins":
        node.append(ref.getparent())
        p._p.append(node)
    else:
        node.append(p._p)
        doc.element.body.insert(0, node)
    note_part(doc, "footnote", [(7, None, "SENTETİK: Not.")])
    with pytest.raises(ValueError, match="^DOCX note reference context is unsupported$"):
        save_and_extract(doc, tmp_path)


@pytest.mark.parametrize("limit", ["text", "work", "depth"])
def test_notes_share_all_body_budgets_and_never_return_partial_passages(tmp_path, monkeypatch, limit):
    doc = Document()
    p = doc.add_paragraph("SENTETİK: Gövde 1234567890.")
    note_reference(p, "footnote", 7)
    blocks = Document()
    blocks.add_paragraph("SENTETİK: Not 1234567890.")
    if limit == "text":
        monkeypatch.setattr("app.extract.MAX_TEXT", 40)
        error = "Extracted text limit exceeded"
    elif limit == "work":
        monkeypatch.setattr("app.extract.MAX_DOCX_BLOCKS", 15)
        for _ in range(20):
            blocks.add_paragraph("")
        error = "DOCX structural work limit exceeded"
    else:
        cell = blocks
        for _ in range(5):
            cell = cell.add_table(rows=1, cols=1).cell(0, 0)
        error = "DOCX table nesting limit exceeded"
    note_part(doc, "footnote", [(7, None, blocks)])
    with pytest.raises(ValueError, match="^" + error + "$"):
        save_and_extract(doc, tmp_path)


def test_note_revision_markup_keeps_accepted_version_unverified(tmp_path):
    doc = Document()
    note_reference(doc.add_paragraph("SENTETİK: Gövde."), "endnote", 7)
    blocks = Document()
    p = blocks.add_paragraph("SENTETİK: Not.")
    p._p.append(OxmlElement("w:ins"))
    note_part(doc, "endnote", [(7, None, blocks)])
    assert "İzlenen değişiklikler var; kabul edilmiş sürüm doğrulanmadı." in save_and_extract(
        doc, tmp_path
    )["warnings"]
