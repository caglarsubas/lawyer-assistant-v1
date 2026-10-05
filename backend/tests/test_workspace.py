import base64
import hashlib
import io
import time
from datetime import datetime, timedelta, timezone

import httpx
import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader
from sqlalchemy import select

from app.auth import hash_password
from app.config import Settings
from app.db import Membership, Record, User
from app.main import create_app
from app.research import run_research, select_passages


@pytest.fixture
def workspace(tmp_path):
    settings = Settings(
        _env_file=None,
        demo_mode=True,
        cookie_secure=False,
        data_dir=tmp_path,
        LLM_PROVIDER_BASE_URL="",
        LLM_PROVIDER_API_KEY="",
        LLM_PROVIDER_MODEL="",
    )
    app = create_app(settings)
    with TestClient(app) as client:
        login = client.post("/api/v1/auth/login", json={"username": "demo", "password": "demo-local-only"})
        assert login.status_code == 200
        client.headers["X-CSRF-Token"] = login.json()["csrf_token"]
        yield app, client, client.get("/api/v1/matters").json()[0]["id"]


def research_product(client, matter):
    base = f"/api/v1/matters/{matter}"
    run = client.post(base + "/research", json={"question": "Ödeme ve sona erme bildirimleri"})
    assert run.status_code == 202, run.text
    for _ in range(100):
        state = client.get(base + "/research/" + run.json()["id"]).json()
        if state["status"] in {"completed", "failed"}:
            break
        time.sleep(0.02)
    assert state["status"] == "completed", state
    return client.get(base + "/products/" + state["product_id"]).json()


def test_login_csrf_and_origin(workspace):
    app, client, matter = workspace
    client.headers.pop("X-CSRF-Token")
    payload = {"title": "Yeni dosya", "domain": "employment"}
    assert client.post("/api/v1/matters", json=payload).status_code == 403
    csrf = client.get("/api/v1/auth/me").json()["csrf_token"]
    assert (
        client.post(
            "/api/v1/matters",
            json=payload,
            headers={"X-CSRF-Token": csrf, "Origin": "https://attacker.invalid"},
        ).status_code
        == 403
    )
    client.cookies.clear()
    assert client.get(f"/api/v1/matters/{matter}").status_code == 401
    assert client.get("/api/v1/status").status_code == 401
    assert set(client.get("/api/v1/bootstrap").json()) == {"demo_mode", "version"}


@pytest.mark.parametrize("firm_id", ["demo-firm", "another-firm"])
def test_matter_and_evidence_isolation_even_same_firm_admin(workspace, firm_id):
    app, client, matter = workspace
    product = research_product(client, matter)
    doc = client.get(f"/api/v1/matters/{matter}").json()["documents"][0]
    with app.state.store.session() as session:
        session.add(
            User(
                username="other",
                name="Other",
                firm_id=firm_id,
                role="admin",
                password_hash=hash_password("different-password"),
            )
        )
        session.commit()
    login = client.post("/api/v1/auth/login", json={"username": "other", "password": "different-password"})
    client.headers["X-CSRF-Token"] = login.json()["csrf_token"]
    assert client.get("/api/v1/matters").json() == []
    for suffix in [
        "",
        "/documents/" + doc["id"],
        "/products/" + product["id"],
        "/products/" + product["id"] + "/export",
    ]:
        assert client.get(f"/api/v1/matters/{matter}" + suffix).status_code == 404
    assert (
        client.post(
            f"/api/v1/matters/{matter}/research", json={"question": "unauthorized request"}
        ).status_code
        == 404
    )


def test_ingestion_encryption_evidence_corrections_and_staleness(workspace):
    app, client, matter = workspace
    base = f"/api/v1/matters/{matter}"
    text = "Gizli müşteri örneği 9417253: Ödeme 30 gün içinde yapılır."
    uploaded = client.post(base + "/documents", files={"file": ("../özel.txt", text.encode(), "text/plain")})
    assert uploaded.status_code == 201, uploaded.text
    assert uploaded.json()["name"] == "özel.txt"
    assert "original_path" not in uploaded.json()
    document = client.get(base + "/documents/" + uploaded.json()["id"]).json()
    evidence = document["passages"][0]["id"]
    assert document["passages"][0]["text"] == text
    assert client.get(base + "/documents/" + document["id"] + "/original").content == text.encode()
    with app.state.store.session() as session:
        for row in session.scalars(select(Record)):
            assert "9417253" not in row.payload
    vault = next((app.state.settings.data_dir / "documents").glob("*.enc"))
    assert text.encode() not in vault.read_bytes()
    assert app.state.store.cipher.decrypt(vault.read_bytes()) == text.encode()
    assert client.post(base + "/facts", json={"text": "X", "status": "documented"}).status_code == 422
    assert (
        client.post(
            base + "/facts", json={"text": "X", "status": "documented", "evidence_id": "foreign"}
        ).status_code
        == 422
    )
    fact = client.post(
        base + "/facts", json={"text": text, "status": "documented", "evidence_id": evidence}
    ).json()
    product = research_product(client, matter)
    cited = {p["id"]: p["text"] for p in product["evidence"]}
    assert all(c["text"] in cited[c["evidence_ids"][0]] for c in product["claims"])
    assert product["status"] == "needs_review"
    assert "legal_corpus_unqualified" in product["critical_gaps"]
    assert (
        client.post(base + "/products/" + product["id"] + "/review", json={"decision": "approve"}).status_code
        == 409
    )
    corrected = client.patch(
        base + "/facts/" + fact["id"],
        json={"text": "İtiraz var", "status": "disputed", "reason": "Müşteri düzeltmesi"},
    ).json()
    assert corrected["history"][0]["text"] == text
    assert client.get(base + "/products/" + product["id"]).json()["status"] == "stale"


def test_exports_preserve_turkish_and_source_locations(workspace):
    _, client, matter = workspace
    product = research_product(client, matter)
    base = f"/api/v1/matters/{matter}/products/{product['id']}/export"
    response = client.get(base + "?format=docx")
    assert response.status_code == 200
    paragraphs = "\n".join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)
    assert "GİZLİ" in paragraphs and "Kaynak:" in paragraphs and "eksikler" in paragraphs
    response = client.get(base + "?format=pdf")
    assert response.status_code == 200
    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages)
    assert "GİZLİ" in text and "Kaynak:" in text and "sözleşme" in text


def test_production_intake_requires_worker_and_never_falls_back(workspace, monkeypatch):
    from app.extraction_client import ExtractionError

    app, client, matter = workspace
    app.state.settings.demo_mode = False
    base = f"/api/v1/matters/{matter}/documents"

    def upload():
        return client.post(base, files={"file": ("synthetic.txt", b"Synthetic fixture", "text/plain")})

    assert upload().status_code == 503  # Required malware control missing.
    app.state.settings.clamav_host = "scanner"
    assert upload().status_code == 503  # Required isolated worker missing.

    class Worker:
        failed = False
        revoke = False

        def execute(self, content, suffix):
            assert content == b"Synthetic fixture" and suffix == ".txt"
            if self.revoke:
                with app.state.store.session() as session:
                    user = session.scalar(select(User).where(User.username == "demo"))
                    user.active = False
                    session.commit()
            if self.failed:
                raise ExtractionError("Unavailable")
            return {
                "passages": [{"text": "Synthetic fixture", "locator": "Paragraf 1"}],
                "warnings": [],
                "page_count": None,
                "passage_count": 1,
            }

    monkeypatch.setattr("app.main.scan_document", lambda *args, **kwargs: None)
    monkeypatch.setattr(
        "app.main.subprocess.run", lambda *args, **kwargs: pytest.fail("Production parsed inside API")
    )
    app.state.extractor = Worker()
    result = upload()
    assert result.status_code == 201 and result.json()["status"] == "needs_review"
    app.state.extractor.failed = True
    result = upload()
    assert result.status_code == 201 and result.json()["status"] == "failed"
    assert result.json()["passage_count"] == 0
    assert result.json()["extraction_warnings"]
    originals_before = set((app.state.settings.data_dir / "documents").glob("*.enc"))
    app.state.extractor.failed = False
    app.state.extractor.revoke = True
    assert upload().status_code == 401
    assert set((app.state.settings.data_dir / "documents").glob("*.enc")) == originals_before


def test_draft_export_pins_immutable_version_and_cannot_substitute_other_record(workspace):
    _, client, matter = workspace
    base = f"/api/v1/matters/{matter}/practice/drafts"
    first = client.post(
        base, json={"title": "Avukat incelemesi", "text": "İlk değerlendirme: varsayım."}
    ).json()
    second = client.post(
        base + "/" + first["id"] + "/versions",
        json={
            "title": "Avukat incelemesi",
            "text": "Düzeltilmiş değerlendirme.",
            "expected_revision": first["revision"],
            "change_note": "Kaynak incelemesinden sonra",
        },
    )
    assert second.status_code == 201
    response = client.get(
        base + "/" + first["id"] + "/export", params={"version_id": first["latest_version_id"]}
    )
    assert response.status_code == 200
    text = "\n".join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)
    assert "İlk değerlendirme" in text and "Düzeltilmiş" not in text and first["latest_version_id"] in text
    response = client.get(base + "/" + first["id"] + "/export?format=pdf")
    assert "Düzeltilmiş" in PdfReader(io.BytesIO(response.content)).pages[0].extract_text()
    foreign = client.post(base, json={"title": "Başka", "text": "Başka taslak"}).json()
    assert (
        client.get(
            base + "/" + first["id"] + "/export", params={"version_id": foreign["latest_version_id"]}
        ).status_code
        == 404
    )
    product = research_product(client, matter)
    assert (
        product["snapshots"]["practice_versions"][first["id"]]["version_id"]
        == second.json()["latest_version_id"]
    )


def test_gateway_approval_is_exact_expiring_and_not_network_side_effect(workspace):
    app, client, matter = workspace
    base = f"/api/v1/matters/{matter}/gateway"
    result = client.post(
        base + "/evaluate",
        json={"query": "iş hukuku", "destination": "https://mevzuat.gov.tr/", "query_type": "matter"},
    )
    assert result.status_code == 200
    assert result.json()["decision"] == "REQUIRE_APPROVAL"
    reference = {"request_id": result.json()["request_id"]}
    assert client.post(base + "/execute", json=reference).status_code == 403
    assert client.post(base + "/approve", json=reference).status_code == 200
    assert client.post(base + "/execute", json=reference).status_code == 503
    with app.state.store.session() as session:
        row = session.get(Record, reference["request_id"])
        data = app.state.store.decode(row)
        data["query"] = "ticaret hukuku"
        app.state.store.update(row, data)
        session.commit()
    assert client.post(base + "/execute", json=reference).status_code == 403
    with app.state.store.session() as session:
        row = session.get(Record, reference["request_id"])
        data = app.state.store.decode(row)
        data["expires_at"] = (datetime.now(timezone.utc) - timedelta(minutes=1)).isoformat()
        app.state.store.update(row, data)
        session.commit()
    assert client.post(base + "/approve", json=reference).status_code == 403


def test_unsupported_format_and_spoofed_pdf_fail_visible(workspace):
    _, client, matter = workspace
    url = f"/api/v1/matters/{matter}/documents"
    assert (
        client.post(url, files={"file": ("unsafe.exe", b"MZ", "application/octet-stream")}).status_code == 415
    )
    assert client.post(url, files={"file": ("fake.pdf", b"not PDF", "application/pdf")}).status_code == 415
    failed = client.post(url, files={"file": ("binary.txt", b"\xff\xfe", "text/plain")})
    assert failed.status_code == 201 and failed.json()["status"] == "failed"
    assert failed.json()["extraction_warnings"]


def test_provider_fabrication_is_not_released(workspace):
    app, client, matter = workspace

    class FabricatingProvider:
        configured = True

        def generate(self, question, evidence):
            return '{"summary":"Invented law","claims":[{"text":"Invented legal result", "evidence_ids":["fake"]}]}'

    app.state.provider = FabricatingProvider()
    base = f"/api/v1/matters/{matter}"
    run = client.post(base + "/research", json={"question": "Ödeme koşulları"}).json()
    for _ in range(100):
        state = client.get(base + "/research/" + run["id"]).json()
        if state["status"] == "failed":
            break
        time.sleep(0.02)
    assert state["status"] == "failed"
    assert client.get(base).json()["products"] == []


def test_long_passage_keeps_exact_source_offsets_and_turkish_relevance():
    text = "Arka plan. " * 2000 + "ÖDEME, teslimden sonra yapılır." + " Son bölüm." * 1000
    selected = select_passages("Ödeme koşulları", [{"id": "source", "text": text}])
    assert len(selected) == 1
    passage = selected[0]
    assert "ÖDEME, teslimden sonra yapılır." in passage["text"]
    assert text[passage["excerpt_start"] : passage["excerpt_end"]] == passage["text"]
    assert passage["id"] == "source"
    assert len(passage["text"].encode()) <= 1800


def test_concurrent_review_cannot_overwrite_newer_stale_state(workspace):
    from sqlalchemy.orm.exc import StaleDataError

    app, client, matter = workspace
    product = research_product(client, matter)
    store = app.state.store
    with store.session() as first, store.session() as second:
        older = first.get(Record, product["id"])
        newer = second.get(Record, product["id"])
        store.update(newer, {**store.decode(newer), "status": "stale"})
        second.commit()
        store.update(older, {**store.decode(older), "status": "reviewed"})
        with pytest.raises(StaleDataError):
            first.commit()
    assert client.get(f"/api/v1/matters/{matter}/products/{product['id']}").json()["status"] == "stale"


@pytest.mark.parametrize("corruption", [None, "source_url", "query", "sha256", "media_type"])
def test_gateway_checks_inbound_binding_forces_quarantine_and_prevents_replay(
    workspace, monkeypatch, corruption
):
    app, client, matter = workspace
    app.state.settings.gateway_enabled = True
    app.state.settings.gateway_url = "http://127.0.0.1:8001"
    app.state.settings.gateway_token = "synthetic-internal-gateway-token"
    raw = b"synthetic public source fixture"
    source = {
        "source_url": "https://mevzuat.gov.tr/",
        "query": "iş hukuku",
        "media_type": "text/plain",
        "sha256": hashlib.sha256(raw).hexdigest(),
        "content_base64": base64.b64encode(raw).decode(),
        "status": "published",
        "rights_status": "permitted",
        "untrusted_extra": "must not survive",
    }
    if corruption:
        source[corruption] = "different"
    client_type = httpx.Client
    monkeypatch.setattr(
        httpx,
        "Client",
        lambda **kwargs: client_type(
            transport=httpx.MockTransport(lambda request: httpx.Response(200, json=source)), **kwargs
        ),
    )
    base = f"/api/v1/matters/{matter}/gateway"
    result = client.post(
        base + "/evaluate",
        json={"query": "iş hukuku", "destination": "https://mevzuat.gov.tr/", "query_type": "public"},
    )
    reference = {"request_id": result.json()["request_id"]}
    response = client.post(base + "/execute", json=reference)
    if corruption:
        assert response.status_code == 502
    else:
        assert response.status_code == 200
        assert response.json()["status"] == "quarantined"
        assert response.json()["rights_status"] == "unverified"
        assert "untrusted_extra" not in response.json()
        assert "content_base64" not in response.json()
    assert client.post(base + "/execute", json=reference).status_code == 403


@pytest.mark.parametrize("change", ["cancel", "revoke"])
def test_search_time_cancellation_or_revocation_prevents_inference(workspace, change):
    app, _, matter = workspace
    store = app.state.store
    with store.session() as session:
        user = session.scalar(select(User).where(User.username == "demo"))
        run = store.add(
            session, "research", user, {"question": "Ödeme koşulları", "status": "queued"}, matter
        )
        session.commit()
    calls = []

    class ChangingSearch:
        def search(self, *args, **kwargs):
            with store.session() as session:
                if change == "cancel":
                    current = session.get(Record, run.id)
                    store.update(current, {**store.decode(current), "status": "cancelled"})
                else:
                    session.delete(session.get(Membership, (matter, user.id)))
                session.commit()
            return {"hits": [], "coverage": {"status": "available"}, "snapshot": {}, "limitations": []}

    class ForbiddenProvider:
        configured = True

        def generate(self, *args):
            calls.append("leaked")
            raise AssertionError("Inference must not receive private documents")

    app.state.search = ChangingSearch()
    app.state.provider = ForbiddenProvider()
    run_research(app, run.id)
    assert calls == []
    with store.session() as session:
        assert store.decode(session.get(Record, run.id))["status"] == (
            "cancelled" if change == "cancel" else "failed"
        )
        assert not session.scalar(select(Record).where(Record.kind == "product"))
