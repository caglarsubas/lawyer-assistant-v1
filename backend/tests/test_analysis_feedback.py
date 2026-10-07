"""Invented reviewer findings and mocked inference; no semantic/legal qualification."""

import copy
import io
import json
from threading import Event
from uuid import uuid4

import httpx
import pytest
from docx import Document
from test_analysis_reviews import review_request
from test_analysis_suggestions import adopt, finished, patch, prepare
from test_analysis_workbench import workbench as workbench_fixture

from app import analysis_reviews
from app.analysis_feedback import editable_targets
from app.analysis_proposals import apply_patch, parse_patch, prompt_measurement, proposal_messages
from app.db import Record, digest
from app.evidence_prompt import canonical
from app.provider import ProviderError


@pytest.fixture
def workbench(tmp_path):
    yield from workbench_fixture.__wrapped__(tmp_path)


def setup_feedback(workbench, monkeypatch, *, incomplete=False, targets=('analysis',)):
    _, client, _, *_ = workbench
    record, endpoint = prepare(workbench, monkeypatch, incomplete=incomplete)
    reviews = endpoint.removesuffix('/suggestions') + '/reviews'
    body, _ = review_request(client, reviews, record, decision='changes_requested')
    body['note'] = 'SYNTHETIC unselected review summary'
    body['findings'] = [{'target_id': target, 'severity': 'major', 'text': f'SYNTHETIC finding {index}',
                         'suggested_change': f'SYNTHETIC requested revision {index}', 'evidence_ids': ['synthetic-clause']}
                        for index, target in enumerate(targets)]
    response = client.post(reviews, json=body)
    assert response.status_code == 201, response.text
    return record, endpoint, reviews, response.json()


def request_feedback(client, endpoint, record, event, *, indices=(0,), mode='single'):
    body = {'expected_revision': record['revision'], 'version_id': record['latest_version_id'],
            'request_id': uuid4().hex, 'mode': mode,
            'review_feedback': {'review_id': event['id'], 'finding_indices': list(indices)}}
    response = client.post(endpoint, json=body)
    assert response.status_code == 202, response.text
    return response.json(), body


def response(finding='finding:0', *, targets=('application:a1', 'conclusion'), outcome='proposed_change', text='SYNTHETIC proposed interpretation'):
    return {'finding_id': finding, 'outcome': outcome, 'edited_targets': list(targets), 'text': text,
            'evidence_ids': ['synthetic-clause']}


def test_opt_in_selected_feedback_is_pinned_minimal_and_adoption_never_resolves_the_review(workbench, monkeypatch):
    app, client, base, payload, *_ = workbench
    record, endpoint, reviews, event = setup_feedback(workbench, monkeypatch, targets=('rule:r1', 'conclusion'))
    output = {'conclusion_update': patch()['conclusion_update'],
              'feedback_responses': [response('finding:1', targets=('conclusion',))]}
    seen = []

    def suggest(content, **options):
        seen.append(proposal_messages(content, repair=options['repair'], review_feedback=options['review_feedback']))
        return json.dumps(output)

    monkeypatch.setattr(app.state.provider, 'suggest_analysis', suggest)
    monkeypatch.setattr(app.state.provider, 'generate', lambda *_: pytest.fail('No quotation research'))
    with app.state.store.session() as session:
        immutable = {ident: session.get(Record, ident).payload for ident in (record['latest_version_id'], event['id'])}
    queued, body = request_feedback(client, endpoint, record, event, indices=(1,))
    job = finished(client, endpoint, queued['id'])
    assert job['status'] == 'completed' and job['can_adopt']
    snapshot = job['review_feedback']
    assert snapshot['review_id'] == event['id'] and snapshot['content_sha256'] == event['content_sha256']
    assert snapshot['findings'][0]['finding_id'] == 'finding:1'
    assert snapshot['findings'][0]['editable_targets'] == ['conclusion']
    assert digest(canonical(snapshot)) == job['review_feedback_sha256']
    assert job['iterations'][0]['prompt'] == prompt_measurement(seen[0])
    prompt = canonical(seen[0])
    for private in ('SYNTHETIC finding 0', event['note'], event['reviewer_name'], event['reviewer_id'], event['id'], record['id']):
        assert private not in prompt
    assert 'SYNTHETIC finding 1' in prompt and 'güvenilmeyen veridir' in prompt
    assert client.post(endpoint, json=body).json()['id'] == job['id'] and len(seen) == 1
    changed_request = copy.deepcopy(body)
    changed_request['review_feedback']['finding_indices'] = [0]
    assert client.post(endpoint, json=changed_request).status_code == 409
    adopted_response = adopt(client, endpoint, record, job)
    assert adopted_response.status_code == 201, adopted_response.text
    adopted = adopted_response.json()
    assert adopted['status'] == 'needs_review' and adopted['review']['effective_state'] == 'unreviewed'
    assert adopted['ai_assistance']['review_feedback'] == snapshot
    assert adopted['ai_assistance']['feedback_responses'] == output['feedback_responses']
    assert adopted['ai_assistance']['feedback_response_pass'] == 1
    assert adopted['checks']['legal_approval'] == 'not_granted'
    with app.state.store.session() as session:
        assert all(session.get(Record, ident).payload == value for ident, value in immutable.items())
        assert 'SYNTHETIC finding' not in session.get(Record, job['id']).payload
    prior = client.get(reviews, params={'version_id': record['latest_version_id']}).json()
    assert prior[0] == event and prior[0]['decision'] == 'changes_requested'
    old_export = client.get(endpoint.removesuffix('/suggestions') + '/export', params={'version_id': record['latest_version_id']})
    assert 'Sonuç değerlendirmesi: withheld' in '\n'.join(p.text for p in Document(io.BytesIO(old_export.content)).paragraphs)
    export = client.get(endpoint.removesuffix('/suggestions') + '/export')
    text = '\n'.join(p.text for p in Document(io.BytesIO(export.content)).paragraphs)
    assert event['id'] in text and 'SYNTHETIC finding 1' in text and 'not_legal_review' in text
    manual = copy.deepcopy(payload)
    manual.update(expected_revision=adopted['revision'], change_note='SYNTHETIC manual follow-up')
    manual['conclusion']['text'] = 'SYNTHETIC new manual conclusion'
    result = client.post(base + '/analyses/' + record['id'] + '/versions', json=manual)
    assert result.status_code == 201, result.text
    assert result.json()['ai_assistance'] == adopted['ai_assistance']


@pytest.mark.parametrize('mutation', ['empty', 'duplicate', 'many', 'bool', 'float', 'negative', 'out_of_range', 'unknown',
                                     'foreign_review', 'prose', 'stale_head', 'reviewed', 'recipe'])
def test_unbound_feedback_is_rejected_before_queue_or_provider(workbench, monkeypatch, mutation):
    app, client, _, *_ = workbench
    record, endpoint, reviews, event = setup_feedback(workbench, monkeypatch)
    body = {'expected_revision': record['revision'], 'version_id': record['latest_version_id'], 'request_id': uuid4().hex,
            'review_feedback': {'review_id': event['id'], 'finding_indices': [0]}}
    if mutation in {'empty', 'duplicate', 'many', 'bool', 'float', 'negative', 'out_of_range', 'unknown'}:
        body['review_feedback']['finding_indices'] = {'empty': [], 'duplicate': [0, 0], 'many': list(range(6)),
                                                     'bool': [True], 'float': [0.0], 'negative': [-1],
                                                     'out_of_range': [20], 'unknown': [1]}[mutation]
    elif mutation == 'foreign_review':
        body['review_feedback']['review_id'] = 'SYNTHETIC foreign review'
    elif mutation == 'prose':
        body['review_feedback']['text'] = 'Caller cannot supply feedback prose'
    elif mutation in {'stale_head', 'reviewed'}:
        next_body, _ = review_request(client, reviews, record, decision='changes_requested' if mutation == 'stale_head' else 'reviewed_conditional')
        assert client.post(reviews, json=next_body).status_code == 201
        if mutation == 'reviewed':
            body['review_feedback']['review_id'] = client.get(reviews, params={'version_id': record['latest_version_id']}).json()[0]['id']
    else:
        monkeypatch.setattr(analysis_reviews, 'RECIPE', 'SYNTHETIC changed review policy')
    monkeypatch.setattr(app.state.provider, 'suggest_analysis', lambda *_a, **_k: pytest.fail('Invalid feedback must not reach inference'))
    monkeypatch.setattr(app.state.research_jobs, 'reserve', lambda: pytest.fail('No queue reservation for invalid feedback'))
    result = client.post(endpoint, json=body)
    assert result.status_code in {409, 422}, result.text
    assert client.get(endpoint).json() == []


def test_review_prose_is_not_shared_without_explicit_finding_selection(workbench, monkeypatch):
    app, client, *_ = workbench
    record, endpoint, _, event = setup_feedback(workbench, monkeypatch)

    def suggest(content, **options):
        assert 'review_feedback' not in options
        messages = canonical(proposal_messages(content))
        assert event['findings'][0]['text'] not in messages and 'selected_review_findings' not in messages
        return json.dumps(patch())

    monkeypatch.setattr(app.state.provider, 'suggest_analysis', suggest)
    queued = client.post(endpoint, json={'expected_revision': record['revision'], 'version_id': record['latest_version_id'], 'request_id': uuid4().hex}).json()
    job = finished(client, endpoint, queued['id'])
    assert job['status'] == 'completed' and 'review_feedback' not in job
    assert job['feedback_responses'] == [] and job['feedback_response_pass'] is None


def unit_content():
    return {'title': 'SYNTHETIC', 'issue': 'SYNTHETIC issue', 'posture': '', 'event_date': None,
            'evidence': [{'evidence_id': 'synthetic-clause', 'passage_sha256': 'a' * 64, 'document_revision': 1, 'start': 0, 'end': 10}],
            'premises': [{'id': 'p1', 'kind': 'assumption', 'fact_id': None, 'fact_revision': None, 'text': 'SYNTHETIC assumption'}],
            'rules': [{'id': 'r1', 'kind': 'contract_clause', 'text': 'SYNTHETIC rule', 'evidence_ids': ['synthetic-clause'],
                       'conditions': [{'id': 'c1', 'text': 'SYNTHETIC condition', 'kind': 'element', 'required_status': 'met'}]}],
            'applications': [{'id': 'a1', 'rule_id': 'r1', 'premise_ids': ['p1'], 'rationale': 'SYNTHETIC original',
                              'assessments': [{'condition_id': 'c1', 'status': 'unknown', 'premise_ids': ['p1']}]}],
            'alternatives': [{'id': 'x1', 'kind': 'search_gap', 'text': 'SYNTHETIC missing public research', 'evidence_ids': []}],
            'conclusion': {'text': 'SYNTHETIC original conclusion', 'requested_disposition': 'conditional', 'application_ids': ['a1'],
                           'alternative_ids': ['x1'], 'uncertainty': [], 'next_step': 'SYNTHETIC inspect'}, 'checks': {'defects': []}}


@pytest.mark.parametrize('mutation', ['missing', 'extra', 'duplicate', 'resolved', 'invented_source', 'duplicate_source',
                                     'invented_target', 'duplicate_target', 'noop', 'manual_edit', 'unexplained_edit', 'fixed_rule'])
def test_response_links_cannot_claim_resolution_or_unbound_edits(mutation):
    content = unit_content()
    feedback = {'findings': [{'finding_id': 'finding:0', 'editable_targets': ['application:a1']}]}
    output = {'application_updates': [{'id': 'a1', 'rationale': 'SYNTHETIC candidate'}],
              'feedback_responses': [response(targets=('application:a1',))]}
    item = output['feedback_responses'][0]
    if mutation == 'missing':
        output.pop('feedback_responses')
    elif mutation == 'extra':
        output['feedback_responses'].append(response('finding:1'))
    elif mutation == 'duplicate':
        output['feedback_responses'].append(item.copy())
    elif mutation == 'resolved':
        item['outcome'] = 'resolved'
    elif mutation == 'invented_source':
        item['evidence_ids'] = ['foreign']
    elif mutation == 'duplicate_source':
        item['evidence_ids'] *= 2
    elif mutation == 'invented_target':
        item['edited_targets'] = ['conclusion']
    elif mutation == 'duplicate_target':
        item['edited_targets'] *= 2
    elif mutation == 'noop':
        output['application_updates'][0]['rationale'] = content['applications'][0]['rationale']
    elif mutation == 'manual_edit':
        item['outcome'] = 'requires_manual_work'
    elif mutation == 'unexplained_edit':
        item.update(outcome='unresolved', edited_targets=[])
    else:
        output['rules'] = []
    with pytest.raises(ValueError):
        apply_patch(content, parse_patch(json.dumps(output)), review_feedback=feedback)
    assert content['applications'][0]['rationale'] == 'SYNTHETIC original'


@pytest.mark.parametrize('target,expected', [('analysis', ['application:a1', 'conclusion']), ('conclusion', ['conclusion']),
                                         ('application:a1', ['application:a1', 'conclusion']), ('premise:p1', ['application:a1', 'conclusion']),
                                         ('rule:r1', ['application:a1', 'conclusion']), ('condition:c1', ['application:a1', 'conclusion']),
                                         ('alternative:x1', ['conclusion']), ('premise:unlinked', [])])
def test_edit_candidates_follow_only_saved_declared_dependencies(target, expected):
    assert editable_targets(unit_content(), target) == expected


@pytest.mark.parametrize('outcome', ['requires_manual_work', 'unresolved'])
def test_manual_or_unresolved_response_never_becomes_an_adoptable_change(workbench, monkeypatch, outcome):
    app, client, *_ = workbench
    record, endpoint, _, event = setup_feedback(workbench, monkeypatch)
    output = {'feedback_responses': [response(targets=(), outcome=outcome)]}
    monkeypatch.setattr(app.state.provider, 'suggest_analysis', lambda *_a, **_k: json.dumps(output))
    queued, _ = request_feedback(client, endpoint, record, event)
    job = finished(client, endpoint, queued['id'])
    assert job['status'] == 'completed' and not job['can_adopt']
    assert job['feedback_responses'] == output['feedback_responses']
    assert adopt(client, endpoint, record, job).status_code == 409


def test_rejected_second_pass_cannot_replace_accepted_candidate_response_lineage(workbench, monkeypatch):
    app, client, *_ = workbench
    record, endpoint, _, event = setup_feedback(workbench, monkeypatch, incomplete=True)
    first = {'application_updates': [{'id': 'a1', 'rationale': 'SYNTHETIC accepted first proposal'}],
             'feedback_responses': [response(targets=('application:a1',), text='SYNTHETIC first response')]}
    rejected = patch()
    rejected['application_updates'][0]['assessments'].append({'condition_id': 'c1', 'status': 'not_met', 'premise_ids': ['p1']})
    rejected['feedback_responses'] = [response(text='SYNTHETIC rejected response')]
    monkeypatch.setattr(app.state.provider, 'suggest_analysis', lambda _content, **options: json.dumps(rejected if options['repair'] else first))
    queued, _ = request_feedback(client, endpoint, record, event, mode='repair')
    job = finished(client, endpoint, queued['id'])
    assert job['status'] == 'completed' and job['can_adopt']
    assert [item['outcome'] for item in job['iterations']] == ['accepted_structural_patch', 'rejected_new_critical_checks']
    assert job['feedback_response_pass'] == 1 and job['feedback_responses'] == first['feedback_responses']
    assert job['iterations'][1]['patch']['feedback_responses'] == rejected['feedback_responses']
    adopted = adopt(client, endpoint, record, job).json()
    assert adopted['ai_assistance']['feedback_responses'] == first['feedback_responses']
    assert adopted['checks']['effective_disposition'] == 'withheld'


@pytest.mark.parametrize('boundary', ['head', 'recipe', 'fact', 'cancel'])
def test_changed_review_or_sources_and_cancellation_block_feedback_publication(workbench, monkeypatch, boundary):
    app, client, base, _, fact, *_ = workbench
    record, endpoint, reviews, event = setup_feedback(workbench, monkeypatch)
    entered, release = Event(), Event()

    def suggest(*_args, **_options):
        entered.set()
        assert release.wait(3)
        return json.dumps({**patch(), 'feedback_responses': [response()]})

    monkeypatch.setattr(app.state.provider, 'suggest_analysis', suggest)
    queued, _ = request_feedback(client, endpoint, record, event)
    try:
        assert entered.wait(2)
        if boundary == 'head':
            body, _ = review_request(client, reviews, record, decision='changes_requested')
            assert client.post(reviews, json=body).status_code == 201
        elif boundary == 'recipe':
            monkeypatch.setattr(analysis_reviews, 'RECIPE', 'SYNTHETIC changed policy')
        elif boundary == 'fact':
            assert client.patch(base + '/facts/' + fact['id'], json={'text': 'SYNTHETIC new allegation', 'status': 'alleged', 'reason': 'SYNTHETIC change'}).status_code == 200
        else:
            assert client.post(endpoint + '/' + queued['id'] + '/cancel').status_code == 200
    finally:
        release.set()
    job = finished(client, endpoint, queued['id'])
    assert job['status'] in {'failed', 'cancelled'} and not job['can_adopt'] and not job.get('candidate')
    assert adopt(client, endpoint, record, {**job, 'candidate_sha256': '0' * 64}).status_code == 409


def test_provider_accounts_exact_selected_feedback_and_rejects_context_overflow(workbench, monkeypatch):
    from test_provider_runtime import relay

    app, client, *_ = workbench
    record, endpoint, _, event = setup_feedback(workbench, monkeypatch)
    monkeypatch.setattr(app.state.provider, 'suggest_analysis', lambda *_a, **_k: json.dumps({**patch(), 'feedback_responses': [response()]}))
    queued, _ = request_feedback(client, endpoint, record, event)
    feedback = finished(client, endpoint, queued['id'])['review_feedback']
    monkeypatch.undo()
    # Configure only this fixture's provider; restore no root environment/secret.
    settings = app.state.settings
    settings.provider_base_url = 'https://10.20.0.8/v1'
    settings.provider_api_key = 'SYNTHETIC-NOT-A-REAL-KEY'
    settings.provider_model = 'SYNTHETIC-local-model'
    settings.provider_identity_verified = settings.provider_cloud_fallback_disabled = True
    monkeypatch.setattr(app.state.provider, 'probe', lambda: {'ready': True})
    seen = []
    client_type = httpx.Client

    def serve(request):
        seen.append(json.loads(request.content))
        return httpx.Response(200, json={'model': settings.provider_model, 'choices': [
            {'finish_reason': 'stop', 'message': {'content': json.dumps({'feedback_responses': [response(targets=(), outcome='unresolved')]})}}]})

    monkeypatch.setattr(httpx, 'Client', lambda **kwargs: client_type(transport=httpx.MockTransport(serve), **kwargs))
    messages = proposal_messages(record, review_feedback=feedback)
    measured = prompt_measurement(messages)
    settings.provider_context_limit = measured['total_upper_bound_units']
    app.state.provider.suggest_analysis(record, review_feedback=feedback)
    assert len(seen) == 1 and seen[0]['messages'] == messages and relay.valid_completion(seen[0], settings.provider_model)
    settings.provider_context_limit -= 1
    with pytest.raises(ProviderError):
        app.state.provider.suggest_analysis(record, review_feedback=feedback)
    assert len(seen) == 1


@pytest.mark.parametrize('boundary', ['head', 'recipe', 'feedback_hash', 'response_pass', 'response_text', 'review_content_hash'])
def test_completed_feedback_revalidates_pins_and_response_lineage_before_adoption(workbench, monkeypatch, boundary):
    app, client, *_ = workbench
    record, endpoint, reviews, event = setup_feedback(workbench, monkeypatch)
    monkeypatch.setattr(app.state.provider, 'suggest_analysis', lambda *_a, **_k: json.dumps({**patch(), 'feedback_responses': [response()]}))
    queued, _ = request_feedback(client, endpoint, record, event)
    job = finished(client, endpoint, queued['id'])
    assert job['can_adopt']
    if boundary == 'head':
        body, _ = review_request(client, reviews, record, decision='changes_requested')
        assert client.post(reviews, json=body).status_code == 201
    elif boundary == 'recipe':
        monkeypatch.setattr(analysis_reviews, 'RECIPE', 'SYNTHETIC new review criteria')
    else:
        with app.state.store.session() as session:
            row = session.get(Record, event['id'] if boundary == 'review_content_hash' else job['id'])
            data = app.state.store.decode(row)
            if boundary == 'feedback_hash':
                data['review_feedback_sha256'] = '0' * 64
            elif boundary == 'response_pass':
                data['feedback_response_pass'] = 2
            elif boundary == 'response_text':
                data['feedback_responses'][0]['text'] = 'SYNTHETIC response not from the retained accepted pass'
            else:
                data['content_sha256'] = '0' * 64
            app.state.store.update(row, data)
            session.commit()
    current = client.get(endpoint + '/' + job['id']).json()
    assert current['freshness']['status'] == 'stale' and not current['can_adopt']
    assert adopt(client, endpoint, record, job).status_code == 409
    assert len(client.get(endpoint.removesuffix('/suggestions') + '/versions').json()) == 1


def test_complete_multiple_finding_responses_allow_shared_actual_edits_without_changing_fixed_inputs():
    content = unit_content()
    feedback = {'findings': [{'finding_id': 'finding:0', 'editable_targets': ['application:a1']},
                             {'finding_id': 'finding:1', 'editable_targets': ['application:a1', 'conclusion']}]}
    output = {'application_updates': [{'id': 'a1', 'rationale': 'SYNTHETIC proposed rationale'}],
              'feedback_responses': [response(targets=('application:a1',)), response('finding:1', targets=('application:a1',))]}
    result = apply_patch(content, parse_patch(json.dumps(output)), review_feedback=feedback)
    assert result.applications[0].rationale == output['application_updates'][0]['rationale']
    for key in ('premises', 'rules', 'alternatives', 'conclusion'):
        assert result.model_dump()[key] == content[key]
    with pytest.raises(ValueError):
        apply_patch(content, parse_patch(json.dumps(output)))
