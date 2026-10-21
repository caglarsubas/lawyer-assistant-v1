"""Exact text-box source contexts and conservative shape/alternative boundaries."""

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx_note_fixtures import note_part, note_reference
from docx_textbox_fixtures import MC, WPS, alternate_textbox, textbox
from lxml import etree

from app.extract import extract


def blocks(text="SENTETİK: Kutudaki istisna."):
    doc = Document()
    doc.add_paragraph(text)
    return doc


def extracted(doc, tmp_path):
    path = tmp_path / "invented.docx"
    doc.save(path)
    original = path.read_bytes()
    result = extract(path, ".docx")
    assert path.read_bytes() == original
    return result


@pytest.mark.parametrize("kind", ["modern", "legacy"])
def test_textbox_is_a_separate_exact_context_without_relocating_body(tmp_path, kind):
    doc = Document()
    p = doc.add_paragraph("SENTETİK: Önce ")
    textbox(p, blocks(), kind)
    p.add_run("sonra.")
    doc.add_paragraph("SENTETİK: Sonraki paragraf.")
    result = extracted(doc, tmp_path)
    assert result["passages"] == [
        {"locator": "Paragraf 1", "text": "SENTETİK: Önce sonra."},
        {"locator": "Paragraf 1, metin kutusu 1, paragraf 1", "text": "SENTETİK: Kutudaki istisna."},
        {"locator": "Paragraf 2", "text": "SENTETİK: Sonraki paragraf."},
    ]
    assert result["extraction_profile"] == "docx-textboxes-v5"
    assert any("görsel okuma sırası" in w for w in result["warnings"])


def test_blank_hosts_multiple_boxes_and_empty_boxes_keep_xml_numbering(tmp_path):
    doc = Document()
    doc.add_paragraph("")
    p = doc.add_paragraph("")
    textbox(p, blocks(""))
    textbox(p, blocks("SENTETİK: İkinci kutu."), "legacy")
    result = extracted(doc, tmp_path)
    assert result["passages"] == [{"locator": "Paragraf 2, metin kutusu 2, paragraf 1",
                                  "text": "SENTETİK: İkinci kutu."}]
    assert any("Boş metin kutusu" in w for w in result["warnings"])


@pytest.mark.parametrize("requires", ["wps", "other"])
def test_equivalent_choice_fallback_is_extracted_once_with_explicit_warning(tmp_path, requires):
    doc = Document()
    alternate_textbox(doc.add_paragraph("SENTETİK: Bağlayıcı."), blocks(), requires=requires)
    result = extracted(doc, tmp_path)
    assert sum(p["text"] == "SENTETİK: Kutudaki istisna." for p in result["passages"]) == 1
    assert any("yinelenen kopya" in w for w in result["warnings"])


@pytest.mark.parametrize("defect", ["different_text", "different_reference", "different_table",
                                    "no_box", "unknown_choice_only", "bad_requires", "duplicate_fallback",
                                    "outside_text", "nested_box", "linked_modern", "linked_legacy",
                                    "revision_wrapper", "content_wrapper", "unknown_parent"])
def test_ambiguous_or_unsupported_shapes_fail_without_partial_success(tmp_path, defect):
    doc = Document()
    p = doc.add_paragraph("SENTETİK: Kısmi başarı dönmemeli.")
    content = blocks()
    if defect in {"different_text", "different_reference", "different_table", "no_box",
                  "unknown_choice_only", "bad_requires", "duplicate_fallback", "outside_text"}:
        other = blocks("SENTETİK: Farklı istisna.") if defect == "different_text" else blocks()
        if defect == "different_reference":
            note_reference(other.paragraphs[0], "footnote", 8)
        if defect == "different_table":
            other.add_table(rows=1, cols=1).cell(0, 0).text = "SENTETİK: Ek koşul."
        alt = alternate_textbox(p, content, other)
        if defect == "no_box":
            alt[1].clear()
        if defect == "unknown_choice_only":
            alt[0].set("Requires", "other")
            alt.remove(alt[1])
        if defect == "bad_requires":
            alt[0].set("Requires", "undefined")
        if defect == "duplicate_fallback":
            alt.append(etree.Element(f"{{{MC}}}Fallback"))
        if defect == "outside_text":
            alt[1].append(OxmlElement("w:t"))
    else:
        if defect == "nested_box":
            textbox(content.paragraphs[0], blocks("SENTETİK: İç kutu."))
        shape, box = textbox(p, content, "legacy" if defect == "linked_legacy" else "modern")
        if defect == "linked_modern":
            shape.append(etree.Element(f"{{{WPS}}}linkedTxbx"))
        if defect == "linked_legacy":
            box.getparent().set("style", "mso-next-textbox:#other")
        if defect in {"revision_wrapper", "content_wrapper"}:
            run = shape.getparent()
            p._p.remove(run)
            wrapper = OxmlElement("w:ins" if defect == "revision_wrapper" else "w:sdt")
            wrapper.append(run)
            p._p.append(wrapper)
        if defect == "unknown_parent":
            box.getparent().tag = qn("w:customXml")
    with pytest.raises(ValueError, match="^DOCX .*text.box.*"):
        extracted(doc, tmp_path)


@pytest.mark.parametrize("story", ["cell", "header", "footer", "footnote", "endnote"])
def test_boxes_retain_table_story_and_note_source_contexts(tmp_path, story):
    doc = Document()
    doc.add_paragraph("SENTETİK: Gövde.")
    note_blocks = None
    if story == "cell":
        p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        prefix = "Tablo 1, satır 1, hücre 1, paragraf 1"
    elif story in {"header", "footer"}:
        p = getattr(doc.sections[0], story).paragraphs[0]
        label = "üstbilgi" if story == "header" else "altbilgi"
        prefix = f"Bölüm 1, {label} tanımı (varsayılan; kaynak bölüm 1), paragraf 1"
    else:
        note_blocks = blocks("SENTETİK: Not açıklaması.")
        p = note_blocks.paragraphs[0]
        note_reference(doc.paragraphs[0], story, 7)
        label = "Dipnot" if story == "footnote" else "Sonnot"
        prefix = f"{label} kimlik 7, başvuru 1 (Paragraf 1), paragraf 1"
    content = blocks()
    content.add_table(rows=1, cols=1).cell(0, 0).text = "SENTETİK: Kutu tablosu."
    textbox(p, content)
    if note_blocks is not None:
        note_part(doc, story, [(7, None, note_blocks)])
    result = extracted(doc, tmp_path)
    assert {"locator": prefix + ", metin kutusu 1, paragraf 1",
            "text": "SENTETİK: Kutudaki istisna."} in result["passages"]
    assert {"locator": prefix + ", metin kutusu 1, tablo 1, satır 1",
            "text": "SENTETİK: Kutu tablosu."} in result["passages"]


def test_note_reference_in_box_is_linked_to_box_not_outer_host(tmp_path):
    doc = Document()
    content = blocks()
    note_reference(content.paragraphs[0], "footnote", 7)
    alternate_textbox(doc.add_paragraph("SENTETİK: Gövde."), content)
    note_part(doc, "footnote", [(7, None, "SENTETİK: Kutunun notu.")])
    result = extracted(doc, tmp_path)
    assert result["passages"][-1] == {
        "locator": "Dipnot kimlik 7, başvuru 1 (Paragraf 1, metin kutusu 1, paragraf 1), paragraf 1",
        "text": "SENTETİK: Kutunun notu.",
    }


def test_separate_equal_boxes_and_inherited_definitions_retain_each_source_context(tmp_path):
    from docx.enum.section import WD_SECTION_START

    doc = Document()
    p = doc.add_paragraph("SENTETİK: Gövde.")
    textbox(p, blocks())
    textbox(p, blocks())
    textbox(doc.sections[0].header.paragraphs[0], blocks())
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    passages = [p for p in extracted(doc, tmp_path)["passages"] if p["text"] == "SENTETİK: Kutudaki istisna."]
    assert len(passages) == 4
    assert len({p["locator"] for p in passages}) == 4
    assert any("Bölüm 2, üstbilgi tanımı (varsayılan; kaynak bölüm 1)" in p["locator"] for p in passages)


def test_box_in_note_cannot_contain_a_nested_note_reference(tmp_path):
    doc = Document()
    content = blocks()
    note_reference(content.paragraphs[0], "endnote", 8)
    parent = blocks()
    textbox(parent.paragraphs[0], content)
    note_reference(doc.add_paragraph("SENTETİK: Gövde."), "footnote", 7)
    note_part(doc, "footnote", [(7, None, parent)])
    with pytest.raises(ValueError, match="^DOCX nested note reference is prohibited$"):
        extracted(doc, tmp_path)


@pytest.mark.parametrize("depth,admitted", [(4, True), (5, False)])
def test_table_depth_is_shared_across_box_boundary(tmp_path, depth, admitted):
    doc = Document()
    outer = doc.add_table(rows=1, cols=1).cell(0, 0)
    content = blocks()
    cell = content.add_table(rows=1, cols=1).cell(0, 0)
    for _ in range(depth - 2):
        cell = cell.add_table(rows=1, cols=1).cell(0, 0)
    cell.text = "SENTETİK: En iç koşul."
    textbox(outer.paragraphs[0], content)
    if admitted:
        assert any(p["text"] == "SENTETİK: En iç koşul." for p in extracted(doc, tmp_path)["passages"])
    else:
        with pytest.raises(ValueError, match="^DOCX table nesting limit exceeded$"):
            extracted(doc, tmp_path)


def test_empty_shape_xml_is_charged_to_the_shared_structural_work_budget(tmp_path, monkeypatch):
    import app.extract as module

    doc = Document()
    shape, _ = textbox(doc.add_paragraph("SENTETİK: Gövde."), blocks(""))
    for _ in range(60):
        shape.append(OxmlElement("w:r"))
    monkeypatch.setattr(module, "MAX_DOCX_BLOCKS", 50)
    with pytest.raises(ValueError, match="^DOCX structural work limit exceeded$"):
        extracted(doc, tmp_path)


def test_unsupported_box_block_wrapper_does_not_drop_its_text(tmp_path):
    doc = Document()
    _, box = textbox(doc.add_paragraph("SENTETİK: Gövde."), blocks())
    node = box[0]
    box.remove(node)
    wrapper = OxmlElement("w:sdt")
    wrapper.append(node)
    box.append(wrapper)
    with pytest.raises(ValueError, match="^DOCX text-box content is unsupported$"):
        extracted(doc, tmp_path)


def test_reference_in_box_hyperlink_has_exact_box_context(tmp_path):
    doc = Document()
    _, box = textbox(doc.add_paragraph("SENTETİK: Gövde."), blocks())
    p = Paragraph(box[0], doc)
    ref = note_reference(p, "footnote", 7)
    run = ref.getparent()
    p._p.remove(run)
    link = OxmlElement("w:hyperlink")
    link.append(run)
    p._p.append(link)
    note_part(doc, "footnote", [(7, None, "SENTETİK: Not.")])
    assert "metin kutusu 1, paragraf 1" in extracted(doc, tmp_path)["passages"][-1]["locator"]
