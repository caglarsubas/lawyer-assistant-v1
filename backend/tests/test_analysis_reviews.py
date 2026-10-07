"""Invented private documents and human declarations; no legal or model approval."""

import copy
import io
import json
from threading import Event
from uuid import uuid4

import pytest
from docx import Document
from sqlalchemy import func, select
from test_analysis_workbench import workbench as workbench_fixture

from app import analysis_reviews, analysis_workbench
from app.analysis_reviews import CRITERIA
from app.db import Audit, Membership, Record


@pytest.fixture
def workbench(tmp_path):
    yield from workbench_fixture.__wrapped__(tmp_path)


def prepare(workbench, payload=None):
    _, client, base, original, *_ = workbench
    created = client.post(base + '/analyses', json=payload or original)
    assert created.status_code == 201, created.text
    record = created.json()
    endpoint = base + '/analyses/' + record['id'] + '/reviews'
    return record, endpoint


def review_request(client, endpoint, record, *, decision='reviewed_conditional'):
    response = client.get(endpoint + '/context', params={'version_id': record['latest_version_id']})
    assert response.status_code == 200, response.text
    context = response.json()
    body = {key: context[key] for key in ('version_id', 'expected_revision', 'content_sha256', 'expected_review_id')}
    body.update(request_id=uuid4().hex, decision=decision, note='SYNTHETIC — Koşullu özel taslak ve sınırları incelendi.',
                criteria=[{'criterion': key, 'outcome': 'confirmed', 'note': ''} for key in CRITERIA], findings=[])
    if decision == 'changes_requested':
        body['findings'] = [{'target_id': 'conclusion', 'severity': 'critical',
                             'text': 'SYNTHETIC — Sonucun özel belge yorumu düzeltilmeli.',
                             'evidence_ids': [record['evidence'][0]['evidence_id']]}]
    return body, context


def test_human_review_projects_only_exact_conditional_draft_and_keeps_immutable_text(workbench, monkeypatch):
    app, client, base, *_ = workbench
    record, endpoint = prepare(workbench)
    monkeypatch.setattr(app.state.provider, 'suggest_analysis', lambda *_a, **_k: pytest.fail('No inference in human review'))
    monkeypatch.setattr(app.state.provider, 'generate', lambda *_a: pytest.fail('No inference in human review'))
    monkeypatch.setattr(app.state.graph, 'tool', lambda *_a: pytest.fail('No authority search in human review'))
    with app.state.store.session() as session:
        frozen = {ident: session.get(Record, ident).payload for ident in (record['id'], record['latest_version_id'])}
    body, context = review_request(client, endpoint, record)
    assert context['can_accept'] and context['review']['effective_state'] == 'unreviewed'
    response = client.post(endpoint, json=body)
    assert response.status_code == 201, response.text
    event = response.json()
    assert event['immutable'] and event['scope'] == 'conditional_private_draft_only'
    assert event['source_selections'][0]['quote_sha256'] == record['evidence'][0]['quote_sha256']
    viewed = client.get(base + '/analyses').json()[0]
    assert viewed['status'] == 'reviewed' and viewed['stored_status'] == 'needs_review'
    assert viewed['checks'] == record['checks'] and viewed['checks']['legal_approval'] == 'not_granted'
    assert viewed['review']['latest']['id'] == event['id']
    assert client.post(endpoint, json=body).json()['id'] == event['id']
    assert client.post(endpoint, json={**body, 'note': 'SYNTHETIC different request'}).status_code == 409
    history = client.get(endpoint, params={'version_id': record['latest_version_id']}).json()
    assert len(history) == 1 and history[0]['id'] == event['id']
    with app.state.store.session() as session:
        assert all(session.get(Record, ident).payload == encrypted for ident, encrypted in frozen.items())
        assert 'SYNTHETIC — Koşullu' not in session.get(Record, event['id']).payload
        assert session.scalar(select(func.count()).select_from(Audit).where(Audit.action == 'analysis_review_recorded')) == 1
    export = client.get(base + '/analyses/' + record['id'] + '/export')
    assert export.status_code == 200
    text = '\n'.join(p.text for p in Document(io.BytesIO(export.content)).paragraphs)
    assert 'Avukat incelemesi: reviewed_conditional' in text and event['content_sha256'] in text
    assert body['note'] in text and 'kamu hukuku otoritesi veya makine hukuki onayı değildir' in text


@pytest.mark.parametrize('mutation', ['hash', 'revision', 'head', 'criteria', 'duplicate', 'needs_change', 'not_applicable',
                                     'critical_finding', 'foreign_target', 'foreign_source', 'extra', 'empty_change'])
def test_unbound_or_incomplete_review_cannot_project_reviewed(workbench, mutation):
    _, client, base, *_ = workbench
    record, endpoint = prepare(workbench)
    body, _ = review_request(client, endpoint, record)
    if mutation == 'hash':
        body['content_sha256'] = '0' * 64
    elif mutation == 'revision':
        body['expected_revision'] += 1
    elif mutation == 'head':
        body['expected_review_id'] = 'SYNTHETIC unknown review'
    elif mutation == 'criteria':
        body['criteria'].pop()
    elif mutation == 'duplicate':
        body['criteria'][1] = body['criteria'][0]
    elif mutation == 'needs_change':
        body['criteria'][0].update(outcome='needs_change', note='SYNTHETIC unresolved source')
    elif mutation == 'not_applicable':
        body['criteria'][0].update(outcome='not_applicable', note='SYNTHETIC cannot skip sources')
    elif mutation in {'critical_finding', 'foreign_target', 'foreign_source'}:
        body['findings'] = [{'target_id': 'conclusion', 'severity': 'note', 'text': 'SYNTHETIC retained limitation', 'evidence_ids': []}]
        if mutation == 'critical_finding':
            body['findings'][0]['severity'] = 'critical'
        elif mutation == 'foreign_target':
            body['findings'][0]['target_id'] = 'foreign'
        else:
            body['findings'][0]['evidence_ids'] = ['foreign']
    elif mutation == 'extra':
        body['legal_approval'] = 'granted'
    elif mutation == 'empty_change':
        body['decision'] = 'changes_requested'
    response = client.post(endpoint, json=body)
    assert response.status_code in {409, 422}, response.text
    assert client.get(base + '/analyses').json()[0]['status'] == 'needs_review'
    assert client.get(endpoint, params={'version_id': record['latest_version_id']}).json() == []


def test_critical_structure_can_be_criticized_but_not_accepted(workbench):
    _, client, base, payload, *_ = workbench
    incomplete = copy.deepcopy(payload)
    incomplete['applications'][0]['assessments'].pop()
    record, endpoint = prepare(workbench, incomplete)
    body, context = review_request(client, endpoint, record)
    assert context['critical_count'] and not context['can_accept']
    assert client.post(endpoint, json=body).status_code == 409
    body, _ = review_request(client, endpoint, record, decision='changes_requested')
    assert client.post(endpoint, json=body).status_code == 201
    viewed = client.get(base + '/analyses').json()[0]
    assert viewed['status'] == 'needs_review' and viewed['review']['effective_state'] == 'changes_requested'
    assert viewed['checks']['effective_disposition'] == 'withheld'


def test_changed_evidence_invalidates_review_without_rewriting_it(workbench):
    app, client, base, _, fact, *_ = workbench
    record, endpoint = prepare(workbench)
    body, _ = review_request(client, endpoint, record)
    accepted = client.post(endpoint, json=body).json()
    with app.state.store.session() as session:
        payload = session.get(Record, accepted['id']).payload
    assert client.patch(base + '/facts/' + fact['id'], json={'text': 'SYNTHETIC new allegation', 'status': 'alleged', 'reason': 'SYNTHETIC corrected role'}).status_code == 200
    stale = client.get(base + '/analyses').json()[0]
    assert stale['status'] == 'stale' and stale['review']['effective_state'] == 'stale'
    export = client.get(base + '/analyses/' + record['id'] + '/export')
    text = '\n'.join(p.text for p in Document(io.BytesIO(export.content)).paragraphs)
    assert 'Güncellik: stale' in text and 'Sonuç değerlendirmesi: withheld' in text
    history = client.get(base + '/analyses/' + record['id'] + '/versions').json()
    assert history[0]['freshness']['status'] == 'stale'
    assert stale['review']['latest']['decision'] == 'reviewed_conditional'
    body, context = review_request(client, endpoint, record)
    assert not context['can_accept'] and client.post(endpoint, json=body).status_code == 409
    with app.state.store.session() as session:
        assert session.get(Record, accepted['id']).payload == payload


def test_new_version_does_not_inherit_review_and_old_history_retains_its_decision(workbench):
    _, client, base, payload, *_ = workbench
    record, endpoint = prepare(workbench)
    body, _ = review_request(client, endpoint, record)
    event = client.post(endpoint, json=body).json()
    revised = client.post(base + '/analyses/' + record['id'] + '/versions', json={
        **payload, 'expected_revision': record['revision'], 'change_note': 'SYNTHETIC explicit correction',
    }).json()
    assert revised['status'] == 'needs_review' and revised['review']['latest'] is None
    body, context = review_request(client, endpoint, record)
    assert not context['current_version'] and not context['can_record']
    assert client.post(endpoint, json=body).status_code == 409
    history = client.get(base + '/analyses/' + record['id'] + '/versions').json()
    assert history[0]['review']['effective_state'] == 'unreviewed'
    assert history[1]['review']['latest']['id'] == event['id']
    assert history[1]['content']['status'] == 'needs_review'


def test_later_change_request_supersedes_projection_without_editing_prior_review(workbench):
    app, client, base, *_ = workbench
    record, endpoint = prepare(workbench)
    body, _ = review_request(client, endpoint, record)
    first = client.post(endpoint, json=body).json()
    with app.state.store.session() as session:
        encrypted = session.get(Record, first['id']).payload
    body, _ = review_request(client, endpoint, record, decision='changes_requested')
    next_review = client.post(endpoint, json=body)
    assert next_review.status_code == 201, next_review.text
    assert next_review.json()['sequence'] == 2 and next_review.json()['expected_review_id'] == first['id']
    assert client.get(base + '/analyses').json()[0]['review']['effective_state'] == 'changes_requested'
    assert client.get(base + '/analyses').json()[0]['status'] == 'needs_review'
    export = client.get(base + '/analyses/' + record['id'] + '/export')
    text = '\n'.join(p.text for p in Document(io.BytesIO(export.content)).paragraphs)
    assert 'Sonuç değerlendirmesi: withheld' in text and 'Avukat değişiklik istedi' in text
    assert client.get(base + '/analyses').json()[0]['checks'] == record['checks']
    with app.state.store.session() as session:
        assert session.get(Record, first['id']).payload == encrypted
    assert len(client.get(endpoint, params={'version_id': record['latest_version_id']}).json()) == 2


def test_review_recipe_change_requires_new_review_without_changing_sources(workbench, monkeypatch):
    _, client, base, *_ = workbench
    record, endpoint = prepare(workbench)
    body, _ = review_request(client, endpoint, record)
    assert client.post(endpoint, json=body).status_code == 201
    monkeypatch.setattr(analysis_reviews, 'RECIPE', 'SYNTHETIC newer review policy')
    stale = client.get(base + '/analyses').json()[0]
    assert stale['status'] == 'stale' and stale['review']['effective_state'] == 'stale'
    export = client.get(base + '/analyses/' + record['id'] + '/export')
    text = '\n'.join(p.text for p in Document(io.BytesIO(export.content)).paragraphs)
    assert 'Güncellik: stale' in text and 'Sonuç değerlendirmesi: withheld' in text
    body, context = review_request(client, endpoint, record)
    assert context['freshness']['status'] == 'current' and context['can_accept']
    assert client.post(endpoint, json=body).status_code == 201
    assert client.get(base + '/analyses').json()[0]['status'] == 'reviewed'


def test_no_ai_criterion_skip_for_ai_adopted_version(workbench, monkeypatch):
    from test_analysis_suggestions import adopt, finished, patch, start
    from test_analysis_suggestions import prepare as suggestion_prepare

    app, client, base, *_ = workbench
    record, endpoint = suggestion_prepare(workbench, monkeypatch)
    monkeypatch.setattr(app.state.provider, 'suggest_analysis', lambda *_a, **_k: json.dumps(patch()))
    queued, _ = start(client, endpoint, record)
    result = finished(client, endpoint, queued['id'])
    adopted = adopt(client, endpoint, record, result).json()
    reviews = base + '/analyses/' + record['id'] + '/reviews'
    body, context = review_request(client, reviews, adopted)
    assert context['ai_contribution_required']
    body['criteria'][-1].update(outcome='not_applicable', note='SYNTHETIC invalid skip')
    assert client.post(reviews, json=body).status_code == 422
    body['criteria'][-1].update(outcome='confirmed', note='SYNTHETIC AI edits inspected')
    assert client.post(reviews, json=body).status_code == 201


def test_finding_targets_disambiguate_nodes_named_like_the_conclusion(workbench):
    _, client, _, payload, *_ = workbench
    payload = copy.deepcopy(payload)
    payload['premises'][0]['id'] = 'conclusion'
    for application in payload['applications']:
        application['premise_ids'] = ['conclusion']
        for item in application['assessments']:
            item['premise_ids'] = ['conclusion']
    record, endpoint = prepare(workbench, payload)
    body, context = review_request(client, endpoint, record, decision='changes_requested')
    assert context['targets'].count('conclusion') == 1 and 'premise:conclusion' in context['targets']
    body['findings'][0]['target_id'] = 'premise:conclusion'
    event = client.post(endpoint, json=body)
    assert event.status_code == 201 and event.json()['findings'][0]['target_id'] == 'premise:conclusion'


def test_review_change_blocks_inflight_proposal_output(workbench, monkeypatch):
    from test_analysis_suggestions import finished, patch, start
    from test_analysis_suggestions import prepare as suggestion_prepare

    app, client, base, *_ = workbench
    record, endpoint = suggestion_prepare(workbench, monkeypatch)
    reviews = base + '/analyses/' + record['id'] + '/reviews'
    entered, release = Event(), Event()

    def suggest(*_a, **_k):
        entered.set()
        assert release.wait(5)
        return json.dumps(patch())

    monkeypatch.setattr(app.state.provider, 'suggest_analysis', suggest)
    try:
        queued, _ = start(client, endpoint, record)
        assert entered.wait(5)
        body, _ = review_request(client, reviews, record, decision='changes_requested')
        assert client.post(reviews, json=body).status_code == 201
    finally:
        release.set()
    job = finished(client, endpoint, queued['id'])
    assert job['status'] == 'failed' and not job['can_adopt'] and 'candidate' not in job
    assert 'Avukat incelemesi değişti' in job['freshness']['reasons'][0]


def test_later_review_blocks_completed_proposal_adoption_and_preserves_job(workbench, monkeypatch):
    from test_analysis_suggestions import adopt, finished, patch, start
    from test_analysis_suggestions import prepare as suggestion_prepare

    app, client, base, *_ = workbench
    record, endpoint = suggestion_prepare(workbench, monkeypatch)
    monkeypatch.setattr(app.state.provider, 'suggest_analysis', lambda *_a, **_k: json.dumps(patch()))
    queued, _ = start(client, endpoint, record)
    result = finished(client, endpoint, queued['id'])
    assert result['can_adopt']
    with app.state.store.session() as session:
        frozen = session.get(Record, result['id']).payload
    reviews = base + '/analyses/' + record['id'] + '/reviews'
    body, _ = review_request(client, reviews, record, decision='changes_requested')
    assert client.post(reviews, json=body).status_code == 201
    result = client.get(endpoint + '/' + result['id']).json()
    assert result['status'] == 'completed' and not result['can_adopt']
    assert result['freshness']['status'] == 'stale'
    assert adopt(client, endpoint, record, result).status_code == 409
    with app.state.store.session() as session:
        assert session.get(Record, result['id']).payload == frozen


def test_review_history_and_context_are_membership_scoped_and_bounded(workbench):
    app, client, base, *_ = workbench
    record, endpoint = prepare(workbench)
    body, _ = review_request(client, endpoint, record)
    event = client.post(endpoint, json=body).json()
    assert client.get(endpoint, params={'version_id': record['latest_version_id'], 'limit': 21}).status_code == 422
    other = client.post('/api/v1/matters', json={'title': 'SYNTHETIC other', 'domain': 'contracts'}).json()['id']
    foreign = f'/api/v1/matters/{other}/analyses/{record["id"]}/reviews'
    assert client.get(foreign, params={'version_id': record['latest_version_id']}).status_code == 404
    assert client.post(foreign, json=body).status_code == 404
    with app.state.store.session() as session:
        owner = session.get(Record, event['id']).owner_id
        session.delete(session.get(Membership, (base.rsplit('/', 1)[1], owner)))
        session.commit()
    assert client.get(endpoint + '/context', params={'version_id': record['latest_version_id']}).status_code == 404
    assert client.post(endpoint, json=body).status_code == 404


def test_review_change_during_export_rejects_obsolete_response(workbench, monkeypatch):
    _, client, base, *_ = workbench
    record, endpoint = prepare(workbench)
    body, _ = review_request(client, endpoint, record)
    assert client.post(endpoint, json=body).status_code == 201
    render = analysis_workbench.render_export

    def change_review(*args, **kwargs):
        response = render(*args, **kwargs)
        body, _ = review_request(client, endpoint, record, decision='changes_requested')
        assert client.post(endpoint, json=body).status_code == 201
        return response

    monkeypatch.setattr(analysis_workbench, 'render_export', change_review)
    response = client.get(base + '/analyses/' + record['id'] + '/export')
    assert response.status_code == 409 and 'avukat incelemesi değişti' in response.text
