"""Invented authorities and mocked inference; no actual Turkish legal qualification."""

import copy
import io
import json
from threading import Event
from uuid import uuid4

import pytest
from docx import Document
from test_analysis_authorities import payload
from test_analysis_suggestions import adopt, finished
from test_authority_findings import request, saved, setup
from test_authority_findings import workbench as workbench_fixture

from app import analysis_workbench, authority_proposals
from app.analysis_proposals import apply_patch, parse_patch, prompt_measurement, proposal_messages
from app.db import Record
from app.evidence_prompt import canonical

workbench = workbench_fixture


def prepare(workbench, monkeypatch):
    app, client, base, *_ = workbench
    endpoint, context, analysis, reader, permit = setup(workbench, monkeypatch)
    body, _ = request(client, endpoint, outcome='needs_change')
    body['note'] = 'SYNTHETIC UNSELECTED OVERALL NOTE'
    for source in body['sources']:
        for item in source['observations']:
            item['note'] = 'SYNTHETIC selected ' + item['dimension']
    review, _ = saved(client, endpoint, body)
    settings = app.state.settings
    settings.provider_base_url = 'https://10.20.0.8/v1'
    settings.provider_api_key = 'SYNTHETIC-NOT-A-REAL-KEY'
    settings.provider_model = 'SYNTHETIC-local-model'
    settings.provider_context_limit = 32768
    settings.provider_identity_verified = True
    settings.provider_cloud_fallback_disabled = True
    monkeypatch.setattr(app.state.provider, 'probe', lambda: {'ready': True})
    proposal = base + '/analyses/' + analysis['id'] + '/suggestions'
    spec = {'expected_revision': analysis['revision'], 'version_id': analysis['latest_version_id'],
            'request_id': uuid4().hex, 'authority_feedback': {'context_id': context['id'],
                'review_id': review['id'], 'review_sha256': review['review_sha256'],
                'findings': [{'source_index': 0, 'dimension': 'history'}]}}
    return proposal, spec, analysis, review, permit, endpoint, reader


def output(*, text='SYNTHETIC public source dependent text', outcome='proposed_change'):
    return {'conclusion_update': {'text': text, 'next_step': 'SYNTHETIC inspect original evidence',
                                 'uncertainty': ['SYNTHETIC historical applicability is unknown']},
        'authority_responses': [{'finding_id': 'authority:0:history', 'outcome': outcome,
            'edited_targets': ['conclusion'], 'text': 'SYNTHETIC model interpretation, not legal approval',
            'authority_ids': ['authority:0']}]}


def complete(workbench, monkeypatch):
    endpoint, body, analysis, review, permit, reviews, reader = prepare(workbench, monkeypatch)
    monkeypatch.setattr(workbench[0].state.provider, 'suggest_analysis', lambda *_a, **_k: json.dumps(output()))
    response = workbench[1].post(endpoint, json=body)
    assert response.status_code == 202, response.text
    job = finished(workbench[1], endpoint, response.json()['id'])
    assert job['status'] == 'completed', job
    return endpoint, body, analysis, review, permit, reviews, reader, job


def test_explicit_selected_original_evidence_minimal_prompt_and_human_adoption(workbench, monkeypatch):
    app, client, base, *_ = workbench
    endpoint, body, analysis, review, permit, reviews, reader = prepare(workbench, monkeypatch)
    immutable = payload(app, review['id'])
    seen = []

    def suggest(content, **options):
        messages = proposal_messages(content, authority_feedback=options['authority_feedback'])
        seen.append(messages)
        return json.dumps(output())

    monkeypatch.setattr(app.state.provider, 'suggest_analysis', suggest)
    response = client.post(endpoint, json=body)
    assert response.status_code == 202, response.text
    job = finished(client, endpoint, response.json()['id'])
    assert job['status'] == 'completed' and job['can_adopt']
    assert job['iterations'][0]['prompt'] == prompt_measurement(seen[0])
    text = canonical(seen[0])
    quote = review['snapshot']['context_snapshot']['manifest']['sources'][0]['evidence']['text']
    assert quote in text and 'SYNTHETIC selected history' in text
    for excluded in ('SYNTHETIC UNSELECTED OVERALL NOTE', 'SYNTHETIC selected adverse', review['id'],
                     review['reviewer_name'], review['reviewer_id'], analysis['id']):
        assert excluded not in text
    assert client.post(endpoint, json=body).json()['id'] == job['id'] and len(seen) == 1
    assert client.post(endpoint, json={**body, 'mode': 'repair'}).status_code == 409
    assert payload(app, review['id']) == immutable
    result = adopt(client, endpoint, analysis, job)
    assert result.status_code == 201, result.text
    adopted = result.json()
    assert adopted['status'] == 'needs_review' and adopted['review']['effective_state'] == 'unreviewed'
    assert adopted['checks']['legal_approval'] == 'not_granted'
    assert adopted['ai_assistance']['authority_responses'] == output()['authority_responses']
    assert adopted['ai_assistance']['authority_response_pass'] == 1
    assert adopted['authority_dependencies'] == [job['authority_feedback']['dependency']]
    export = client.get(base + '/analyses/' + analysis['id'] + '/export')
    assert export.status_code == 200
    rendered = '\n'.join(p.text for p in Document(io.BytesIO(export.content)).paragraphs)
    assert quote in rendered and review['id'] in rendered
    assert payload(app, review['id']) == immutable


def test_later_private_model_revision_retains_original_public_quotes_and_model_provenance(workbench, monkeypatch):
    app, client, base, *_ = workbench
    endpoint, _, analysis, review, permit, *_rest, job = complete(workbench, monkeypatch)
    first = adopt(client, endpoint, analysis, job).json()
    original = first['authority_contributions']
    app.state.settings.provider_model = 'SYNTHETIC-second-model'
    monkeypatch.setattr(app.state.provider, 'suggest_analysis', lambda *_a, **_k: json.dumps({
        'conclusion_update': {'text': 'SYNTHETIC later private interpretation',
                             'uncertainty': [], 'next_step': 'SYNTHETIC inspect all original sources'}}))
    started = client.post(endpoint, json={'expected_revision': first['revision'],
        'version_id': first['latest_version_id'], 'request_id': uuid4().hex})
    assert started.status_code == 202, started.text
    later = finished(client, endpoint, started.json()['id'])
    result = adopt(client, endpoint, first, later)
    assert result.status_code == 201, result.text
    second = result.json()
    assert second['authority_contributions'] == original
    assert second['ai_assistance']['provider']['model'] == 'SYNTHETIC-second-model'
    assert second['authority_contributions'][0]['provider_pin']['model'] == 'SYNTHETIC-local-model'
    export = client.get(base + '/analyses/' + analysis['id'] + '/export')
    assert export.status_code == 200
    text = '\n'.join(p.text for p in Document(io.BytesIO(export.content)).paragraphs)
    assert review['snapshot']['context_snapshot']['manifest']['sources'][0]['evidence']['text'] in text
    assert job['id'] in text and 'SYNTHETIC-local-model' in text and 'SYNTHETIC-second-model' in text
    permit.active = False
    assert client.get(base + '/analyses/' + analysis['id'] + '/export').status_code == 409
    assert client.get(endpoint + '/' + later['id']).json()['candidate'] is None


def test_denial_culls_jobs_and_adopted_draft_history_exports_and_private_followup(workbench, monkeypatch):
    app, client, base, manual, *_ = workbench
    endpoint, _, analysis, review, permit, _, _, job = complete(workbench, monkeypatch)
    adopted = adopt(client, endpoint, analysis, job).json()
    immutable = payload(app, adopted['latest_version_id'])
    # Subsequent human revisions cannot silently strip the public contribution.
    body = copy.deepcopy(manual)
    body.update(expected_revision=adopted['revision'], change_note='SYNTHETIC human refinement')
    body['conclusion']['text'] = 'SYNTHETIC manual public-dependent revision'
    revision = client.post(base + '/analyses/' + analysis['id'] + '/versions', json=body)
    assert revision.status_code == 201, revision.text
    assert revision.json()['authority_dependencies'] == adopted['authority_dependencies']
    permit.active = False
    for url in (endpoint, endpoint + '/' + job['id'], base + '/research/' + job['id']):
        response = client.get(url)
        assert response.status_code == 200
        assert 'SYNTHETIC public source dependent text' not in response.text
        assert 'SYNTHETIC selected history' not in response.text
        values = response.json() if isinstance(response.json(), list) else [response.json()]
        assert all(item['candidate'] is None and not item['can_adopt'] for item in values)
    for url in (base + '/analyses', base + '/analyses/' + analysis['id'] + '/versions',
                base + '/analyses/' + analysis['id'] + '/export'):
        response = client.get(url)
        assert response.status_code == 409, response.text
        assert 'SYNTHETIC manual public-dependent revision' not in response.text
    assert adopt(client, endpoint, analysis, job).status_code == 409
    current = revision.json()
    assert client.post(endpoint, json={'expected_revision': current['revision'], 'version_id': current['latest_version_id'],
        'request_id': uuid4().hex}).status_code == 409
    assert payload(app, adopted['latest_version_id']) == immutable


@pytest.mark.parametrize('missing', ['authority_dependencies', 'authority_contributions'])
def test_missing_public_dependencies_fail_closed_instead_of_exposing_derived_text(workbench, monkeypatch, missing):
    app, client, base, *_ = workbench
    endpoint, _, analysis, *_rest, job = complete(workbench, monkeypatch)
    adopted = adopt(client, endpoint, analysis, job).json()
    with app.state.store.session() as session:
        row = session.get(Record, adopted['latest_version_id'])
        value = app.state.store.decode(row)
        value['content'].pop(missing)
        app.state.store.update(row, value)
        session.commit()
    response = client.get(base + '/analyses/' + analysis['id'] + '/export')
    assert response.status_code == 409
    assert 'SYNTHETIC public source dependent text' not in response.text


@pytest.mark.parametrize('failure', ['public', 'guard_exit', 'source', 'new_review', 'policy'])
def test_changed_authority_bindings_prevent_adoption_without_rewriting(workbench, monkeypatch, failure):
    app, client, *_ = workbench
    endpoint, _, analysis, review, permit, reviews, reader, job = complete(workbench, monkeypatch)
    before = payload(app, job['id'])
    if failure == 'public':
        permit.active = False
    elif failure == 'guard_exit':
        permit.fail_exit = True
    elif failure == 'source':
        reader.info['pointer']['sequence'] += 1
    elif failure == 'new_review':
        saved(client, reviews)
    else:
        monkeypatch.setattr(authority_proposals, 'RECIPE', 'SYNTHETIC new authority recipe')
    response = client.get(endpoint + '/' + job['id'])
    assert response.status_code == 200 and not response.json()['can_adopt']
    if failure in {'public', 'guard_exit', 'source', 'policy'}:
        assert response.json()['candidate'] is None
    assert adopt(client, endpoint, analysis, job).status_code == 409
    assert payload(app, job['id']) == before


def test_inflight_revocation_discards_output_and_cancellation_never_adopts(workbench, monkeypatch):
    _, client, *_ = workbench
    endpoint, body, analysis, _, permit, *_ = prepare(workbench, monkeypatch)
    entered, release = Event(), Event()

    def blocked(*_a, **_k):
        entered.set()
        assert release.wait(5)
        return json.dumps(output())

    monkeypatch.setattr(workbench[0].state.provider, 'suggest_analysis', blocked)
    try:
        response = client.post(endpoint, json=body)
        assert response.status_code == 202
        assert entered.wait(5)
        permit.active = False
        release.set()
        job = finished(client, endpoint, response.json()['id'])
        assert job['status'] == 'failed' and job['candidate'] is None and not job['can_adopt']
    finally:
        release.set()


def test_adoption_commit_then_late_guard_failure_retains_pending_version(workbench, monkeypatch):
    app, client, base, *_ = workbench
    endpoint, _, analysis, _, permit, *_rest, job = complete(workbench, monkeypatch)
    original = app.state.store.update

    def late(row, data):
        result = original(row, data)
        if row.kind == 'practice_analysis':
            permit.fail_exit = True
        return result

    monkeypatch.setattr(app.state.store, 'update', late)
    response = adopt(client, endpoint, analysis, job)
    assert response.status_code == 409
    permit.fail_exit = False
    # A transient recovery cannot turn an unadmitted commit into visible content.
    assert client.get(base + '/analyses').status_code == 409
    assert client.get(base + '/analyses/' + analysis['id'] + '/versions').status_code == 409
    assert client.get(base + '/analyses/' + analysis['id'] + '/export').status_code == 409


def test_export_time_revocation_discards_rendered_bytes(workbench, monkeypatch):
    _, client, base, *_ = workbench
    endpoint, _, analysis, _, permit, *_rest, job = complete(workbench, monkeypatch)
    assert adopt(client, endpoint, analysis, job).status_code == 201
    render = analysis_workbench.render_export

    def revoke(*args):
        response = render(*args)
        permit.active = False
        return response

    monkeypatch.setattr(analysis_workbench, 'render_export', revoke)
    response = client.get(base + '/analyses/' + analysis['id'] + '/export')
    assert response.status_code == 409 and 'SYNTHETIC public source dependent text' not in response.text


@pytest.mark.parametrize('mutation', ['duplicate', 'bool', 'too_many', 'dimension', 'prose', 'mixed'])
def test_invalid_selections_are_rejected_without_provider_or_reservation(mutation):
    from app.analysis_suggestions import SuggestionInput

    body = {'expected_revision': 1, 'version_id': 'v1', 'request_id': 'a' * 32,
            'authority_feedback': {'context_id': 'c1', 'review_id': 'r1', 'review_sha256': 'b' * 64,
                                   'findings': [{'source_index': 0, 'dimension': 'history'}]}}
    selection = body['authority_feedback']
    if mutation == 'duplicate':
        selection['findings'] *= 2
    elif mutation == 'bool':
        selection['findings'][0]['source_index'] = True
    elif mutation == 'too_many':
        selection['findings'] = [{'source_index': index, 'dimension': 'history'} for index in range(6)]
    elif mutation == 'dimension':
        selection['findings'][0]['dimension'] = 'fabricated'
    elif mutation == 'prose':
        selection['note'] = 'caller supplied public text'
    else:
        body['review_feedback'] = {'review_id': 'r1', 'finding_indices': [0]}
    with pytest.raises(ValueError):
        SuggestionInput.model_validate(body)


@pytest.mark.parametrize('mutation', ['missing', 'invented_authority', 'false_edit', 'resolution', 'unexplained', 'fixed_norm'])
def test_model_cannot_invent_authority_resolve_findings_or_escape_declared_targets(mutation):
    from test_analysis_feedback import unit_content

    content = unit_content()
    feedback = {'findings': [{'finding_id': 'authority:0:history', 'authority_id': 'authority:0',
                             'editable_targets': ['conclusion']}], 'sources': [{'id': 'authority:0'}]}
    value = output()
    if mutation == 'missing':
        value['authority_responses'] = []
    elif mutation == 'invented_authority':
        value['authority_responses'][0]['authority_ids'] = ['authority:1']
    elif mutation == 'false_edit':
        value['authority_responses'][0]['edited_targets'] = ['application:a1']
    elif mutation == 'resolution':
        value['authority_responses'][0]['outcome'] = 'resolved'
    elif mutation == 'unexplained':
        value['authority_responses'][0].update(outcome='unresolved', edited_targets=[])
    else:
        value['rules'] = []
    with pytest.raises(ValueError):
        apply_patch(content, parse_patch(json.dumps(value)), authority_feedback=feedback)


def test_invalid_bound_review_and_supported_selection_fail_before_admission(workbench, monkeypatch):
    app, client, *_ = workbench
    endpoint, body, _, _, _, reviews, _ = prepare(workbench, monkeypatch)
    monkeypatch.setattr(app.state.research_jobs, 'reserve', lambda: pytest.fail('Invalid selection cannot reserve'))
    variants = []
    for key, value in [('review_sha256', 'f' * 64), ('review_id', 'missing'), ('context_id', 'missing')]:
        changed = copy.deepcopy(body)
        changed['authority_feedback'][key] = value
        variants.append(changed)
    changed = copy.deepcopy(body)
    changed['authority_feedback']['findings'][0]['source_index'] = 7
    variants.append(changed)
    for invalid in variants:
        assert client.post(endpoint, json=invalid).status_code in {404, 409, 422}
    assert client.get(endpoint).json() == []
    # A new review is supported, not an instruction to alter the authority.
    review_body, _ = request(client, reviews, outcome='supported')
    review, _ = saved(client, reviews, review_body)
    body['authority_feedback'].update(review_id=review['id'], review_sha256=review['review_sha256'])
    assert client.post(endpoint, json=body).status_code == 422


def test_late_candidate_guard_exit_retains_pending_bytes_even_after_recovery(workbench, monkeypatch):
    app, client, *_ = workbench
    endpoint, body, _, _, permit, *_ = prepare(workbench, monkeypatch)
    original = app.state.store.update

    def fail_after_candidate(row, data):
        result = original(row, data)
        if row.kind == 'research' and data.get('candidate'):
            permit.fail_exit = True
        return result

    monkeypatch.setattr(app.state.store, 'update', fail_after_candidate)
    monkeypatch.setattr(app.state.provider, 'suggest_analysis', lambda *_a, **_k: json.dumps(output()))
    response = client.post(endpoint, json=body)
    assert response.status_code == 202
    job = finished(client, endpoint, response.json()['id'])
    assert job['status'] == 'failed' and job['candidate'] is None
    permit.fail_exit = False
    recovered = client.get(endpoint + '/' + job['id']).json()
    assert recovered['candidate'] is None and not recovered['can_adopt']
    with app.state.store.session() as session:
        retained = app.state.store.decode(session.get(Record, job['id']))
        assert retained['candidate'] and retained['authority_publication_pending']


def test_cancellation_of_source_bound_call_blocks_adoption(workbench, monkeypatch):
    app, client, *_ = workbench
    endpoint, body, analysis, _, _, *_ = prepare(workbench, monkeypatch)
    entered, release = Event(), Event()

    def blocked(*_a, **_k):
        entered.set()
        assert release.wait(5)
        return json.dumps(output())

    monkeypatch.setattr(app.state.provider, 'suggest_analysis', blocked)
    try:
        queued = client.post(endpoint, json=body).json()
        assert entered.wait(5)
        stopped = client.post(endpoint + '/' + queued['id'] + '/cancel')
        assert stopped.status_code == 200 and stopped.json()['status'] == 'cancelling'
        release.set()
        job = finished(client, endpoint, queued['id'])
        assert job['status'] == 'cancelled' and not job['can_adopt']
        assert client.post(endpoint + '/' + job['id'] + '/adopt', json={
            'expected_revision': analysis['revision'], 'candidate_sha256': 'a' * 64,
            'change_note': 'SYNTHETIC should be rejected'}).status_code == 409
    finally:
        release.set()


def test_rejected_repair_preserves_first_authority_response_lineage(workbench, monkeypatch):
    workbench[3]['applications'][0]['assessments'].pop()
    app, client, *_ = workbench
    endpoint, body, analysis, _, _, *_ = prepare(workbench, monkeypatch)
    body['mode'] = 'repair'
    first = output(text='SYNTHETIC first accepted public edit')
    rejected = output(text='SYNTHETIC rejected public edit')
    rejected['application_updates'] = [{'id': 'a1', 'assessments': [
        {'condition_id': 'c1', 'status': 'not_met', 'premise_ids': ['p1']}]}]
    rejected['authority_responses'][0]['edited_targets'].append('application:a1')
    monkeypatch.setattr(app.state.provider, 'suggest_analysis', lambda _content, **options:
        json.dumps(rejected if options['repair'] else first))
    queued = client.post(endpoint, json=body)
    assert queued.status_code == 202, queued.text
    job = finished(client, endpoint, queued.json()['id'])
    assert job['status'] == 'completed' and job['can_adopt'], job
    assert [item['outcome'] for item in job['iterations']] == ['accepted_structural_patch', 'rejected_new_critical_checks']
    assert job['authority_response_pass'] == 1 and job['authority_responses'] == first['authority_responses']
    assert job['candidate']['conclusion']['text'] == first['conclusion_update']['text']
    assert job['iterations'][1]['patch']['authority_responses'] == rejected['authority_responses']
    adopted = adopt(client, endpoint, analysis, job)
    assert adopted.status_code == 201, adopted.text
    assert adopted.json()['ai_assistance']['authority_response_pass'] == 1
