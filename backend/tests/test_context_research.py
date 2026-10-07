"""Synthetic research/API/provider/export integration; no real keys or network."""

import hashlib
import io
import json

import httpx
import pytest
from docx import Document
from pypdf import PdfReader
from sqlalchemy import select
from test_workspace import research_product
from test_workspace import workspace as workspace_fixture

from app.db import Record
from app.evidence_prompt import canonical


@pytest.fixture
def workspace(tmp_path):
    yield from workspace_fixture.__wrapped__(tmp_path)


def lightweight_graph(monkeypatch, app):
    monkeypatch.setattr(app.state.graph, "tool", lambda *args, **kwargs: {
        "snapshot": {"mode": "synthetic"}, "paths": []})


def seed_document(app, matter, text):
    with app.state.store.session() as session:
        record = session.scalar(select(Record).where(Record.kind == "document", Record.matter_id == matter))
        data = app.state.store.decode(record)
        data["passages"] = [{"id": "synthetic-context", "text": text, "locator": "SYNTHETIC line"}]
        app.state.store.update(record, data)
        session.commit()


def provider_transport(app, monkeypatch, *, wrong_quote=False):
    settings = app.state.settings
    settings.provider_base_url = "https://10.20.0.8/v1"
    settings.provider_api_key = "synthetic-context-key"
    settings.provider_model = "synthetic-local-model"
    settings.provider_identity_verified = True
    settings.provider_cloud_fallback_disabled = True
    original, seen = httpx.Client, []

    def serve(request):
        seen.append(request)
        if request.method == "GET":
            return httpx.Response(200, json={"data": [{"id": settings.provider_model,
                "backend": "ollama_http", "format": "ollama_http",
                "request_key_source": "local-inference", "max_model_len": 8192}]})
        payload = json.loads(request.content)
        evidence = json.loads(payload["messages"][1]["content"])["evidence"][0]
        answer = {"summary": "UNSUPPORTED MODEL PROSE", "claims": [{
            "text": "NOT IN SELECTED WINDOW" if wrong_quote else evidence["text"],
            "evidence_ids": [evidence["id"]]}]}
        return httpx.Response(200, json={"model": settings.provider_model,
            "choices": [{"finish_reason": "stop", "message": {"content": json.dumps(answer)}}]})

    monkeypatch.setattr(httpx, "Client", lambda **kwargs: original(
        transport=httpx.MockTransport(serve), **kwargs))
    return seen


def test_actual_serialized_model_envelope_matches_retained_record_and_exports(workspace, monkeypatch):
    app, client, matter = workspace
    lightweight_graph(monkeypatch, app)
    text = "😀 Bilgi.\r\n" * 300 + "Ödeme yapılmaz. Ancak teslim kanıtlanırsa ödeme gerekir.\r\n" + "Son bilgi.\r\n" * 300
    seed_document(app, matter, text)
    seen = provider_transport(app, monkeypatch)
    product = research_product(client, matter)
    assert len(seen) == 2
    payload = json.loads(seen[-1].content)
    digest = hashlib.sha256(canonical(payload["messages"]).encode()).hexdigest()
    pack = product["context_pack"]
    assert pack["prompt"]["messages_sha256"] == digest
    assert product["snapshots"]["context_pack"]["messages_sha256"] == digest
    assert pack["provider_use"] == "validated_quote_response"
    assert pack["inventory"]["shortened_passages"] == 1
    record, selected = pack["selected"][0], product["evidence"][0]
    assert text[record["excerpt_start"]:record["excerpt_end"]] == selected["text"]
    assert product["claims"][0]["text"] == selected["text"]
    assert "UNSUPPORTED MODEL PROSE" not in json.dumps(product)
    transmitted = json.loads(payload["messages"][1]["content"])["evidence"][0]
    assert set(transmitted) == {"id", "text", "partial"} and transmitted["partial"]
    assert "document_name" not in canonical(payload["messages"])
    export = f"/api/v1/matters/{matter}/products/{product['id']}/export"
    response = client.get(export + "?format=docx")
    assert response.status_code == 200
    docx_text = "\n".join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)
    assert "Belge bağlamı ve seçim sınırları" in docx_text and digest in docx_text
    assert f"[{record['excerpt_start']}, {record['excerpt_end']})" in docx_text
    response = client.get(export + "?format=pdf")
    assert response.status_code == 200
    pdf_text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)
    assert "Belge bağlamı ve seçim sınırları" in pdf_text and digest in pdf_text


def test_unconfigured_preparation_keeps_full_selected_quote_without_second_hidden_slice(workspace, monkeypatch):
    app, client, matter = workspace
    lightweight_graph(monkeypatch, app)
    text = "Ödeme. " + "x " * 740
    seed_document(app, matter, text)
    product = research_product(client, matter)
    assert len(text) > 1200 and len(text.encode()) < 1800
    assert product["claims"][0]["text"] == product["evidence"][0]["text"] == text
    assert product["context_pack"]["provider_use"] == "not_configured"
    assert product["status"] == "needs_review"
    assert product["critical_gaps"] == ["legal_corpus_unqualified"]


def test_legacy_product_without_manifest_or_excerpt_ranges_remains_readable_and_exportable(workspace, monkeypatch):
    app, client, matter = workspace
    lightweight_graph(monkeypatch, app)
    product = research_product(client, matter)
    with app.state.store.session() as session:
        record = session.get(Record, product["id"])
        data = app.state.store.decode(record)
        data.pop("context_pack")
        data["snapshots"].pop("context_pack")
        for passage in data["evidence"]:
            for field in ["excerpt_start", "excerpt_end", "full_passage_length"]:
                passage.pop(field)
        app.state.store.update(record, data)
        session.commit()
    base = f"/api/v1/matters/{matter}/products/{product['id']}"
    retained = client.get(base)
    assert retained.status_code == 200 and "context_pack" not in retained.json()
    response = client.get(base + "/export?format=docx")
    assert response.status_code == 200
    text = "\n".join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)
    assert "Kaynak:" in text and "Belge bağlamı ve seçim sınırları" not in text


def test_unfitting_question_prevents_even_metadata_probe_and_returns_coverage_gap(workspace, monkeypatch):
    app, client, matter = workspace
    lightweight_graph(monkeypatch, app)
    seen = provider_transport(app, monkeypatch)
    app.state.settings.provider_context_limit = 1000
    product = research_product(client, matter)
    assert seen == []
    assert not product["claims"]
    assert product["context_pack"]["provider_use"] == "no_selected_evidence"
    assert product["context_pack"]["omission_counts"]["question_exceeds_context"] > 0
    assert any("model çağrısı yapılmadı" in gap for gap in product["coverage"]["gaps"])


def test_unselected_text_cannot_be_published_as_a_claim(workspace, monkeypatch):
    from test_research_jobs import seed_run, state

    from app.research import run_research

    app, _, matter = workspace
    lightweight_graph(monkeypatch, app)
    provider_transport(app, monkeypatch, wrong_quote=True)
    ident = seed_run(app, matter)
    run_research(app, ident)
    run = state(app, ident)
    assert run["status"] == "failed"
    assert "NOT IN SELECTED WINDOW" not in str(run)
    with app.state.store.session() as session:
        assert session.scalar(select(Record).where(Record.kind == "product", Record.matter_id == matter)) is None
