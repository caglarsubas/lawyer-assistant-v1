"""Local text-only exports; markup from sources is never executable."""

import io
from pathlib import Path
from xml.sax.saxutils import escape

from fastapi.responses import StreamingResponse


def render_export(lines: list[str], file_id: str, format: str):
    output = io.BytesIO()
    if format == "docx":
        from docx import Document

        doc = Document()
        doc.add_heading(lines[0], 0)
        for line in lines[1:]:
            doc.add_paragraph(line)
        doc.save(output)
        media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif format == "pdf":
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        styles = getSampleStyleSheet()
        pdfmetrics.registerFont(
            TTFont("NotoSans", str(Path(__file__).parent / "assets" / "NotoSans-Regular.ttf"))
        )
        styles["Normal"].fontName = "NotoSans"
        styles["Normal"].fontSize = 10
        styles["Normal"].leading = 15
        SimpleDocTemplate(output).build(
            [
                element
                for line in lines
                for element in (Paragraph(escape(line), styles["Normal"]), Spacer(1, 8))
            ]
        )
        media = "application/pdf"
    else:
        raise ValueError("Unsupported export format")
    output.seek(0)
    return StreamingResponse(
        output,
        media_type=media,
        headers={"Content-Disposition": f'attachment; filename="preparation-{file_id}.{format}"'},
    )
