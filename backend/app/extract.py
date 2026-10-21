"""Local bounded extraction worker. Container network isolation is required in production."""

import csv
import io
import json
import re
import resource
import subprocess
import sys
import tempfile
import zipfile
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from pathlib import Path

SUPPORTED_SUFFIXES = frozenset(
    {
        ".txt",
        ".md",
        ".csv",
        ".docx",
        ".xlsx",
        ".pdf",
        ".png",
        ".jpg",
        ".jpeg",
        ".tiff",
        ".tif",
        ".bmp",
        ".webp",
        ".gif",
        ".pptx",
        ".odt",
        ".ods",
        ".odp",
        ".html",
        ".htm",
        ".eml",
        ".rtf",
        ".xls",
        ".msg",
        ".doc",
    }
)
ZIP_SUFFIXES = frozenset({".docx", ".xlsx", ".pptx", ".odt", ".ods", ".odp"})
MAX_TEXT = 2_000_000
DOCX_BODY_PROFILE = "docx-textboxes-v5"
MAX_DOCX_TABLE_DEPTH = 4
MAX_DOCX_BLOCKS = 50_000


class PlainHTML(HTMLParser):
    """Extract text without a browser, resource fetches or executable markup."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.skipped = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "template", "iframe", "object", "svg"}:
            self.skipped += 1
        elif tag in {"p", "div", "br", "tr", "li", "h1", "h2", "h3"} and not self.skipped:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "template", "iframe", "object", "svg"}:
            self.skipped = max(0, self.skipped - 1)

    def handle_data(self, data):
        if not self.skipped:
            self.parts.append(data)


def html_text(value):
    parser = PlainHTML()
    parser.feed(value)
    return "".join(parser.parts)


def check_archive(path):
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > 5000 or sum(m.file_size for m in members) > 100 * 1024 * 1024:
            raise ValueError("Archive expansion limit exceeded")
        if len({m.filename for m in members}) != len(members):
            raise ValueError("Duplicate archive members")
        if any("vbaproject" in m.filename.casefold() or m.filename.startswith("Basic/") for m in members):
            raise ValueError("Macros are prohibited")


def extract(path, suffix):
    passages, warnings = [], []
    page_count = None
    text_size = 0
    if suffix in ZIP_SUFFIXES:
        check_archive(path)

    def add(locator, text):
        nonlocal text_size
        text_size += len(text)
        if text_size > MAX_TEXT or len(passages) >= 100000:
            raise ValueError("Extracted text limit exceeded")
        if text.strip():
            passages.append({"locator": locator, "text": text.strip()})

    if suffix in (".txt", ".md", ".csv"):
        text = path.read_text(encoding="utf-8-sig", errors="strict")
        if suffix == ".csv":
            for index, row in enumerate(csv.reader(io.StringIO(text)), 1):
                add(f"Satır {index}", " | ".join(row))
        else:
            for index, paragraph in enumerate(text.split("\n\n"), 1):
                add(f"Paragraf {index}", paragraph)
    elif suffix in (".docx", ".xlsx"):
        if suffix == ".docx":
            from types import SimpleNamespace

            from defusedxml.ElementTree import fromstring
            from docx import Document
            from docx.blkcntnr import BlockItemContainer
            from docx.opc.constants import CONTENT_TYPE as CT
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
            from docx.oxml import parse_xml
            from docx.oxml.ns import qn
            from docx.parts.hdrftr import FooterPart, HeaderPart
            from docx.table import Table, _Cell
            from docx.text.paragraph import Paragraph

            from app.docx_textboxes import BOX, prepare_textboxes, validate_box_context

            doc = Document(path)
            paragraph_index = table_index = 0
            visited = 0
            has_tables = False
            references = []
            observed_references = set()
            source_roots = [doc.element]
            note_tags = {qn("w:" + kind + "Reference"): kind for kind in ("footnote", "endnote")}
            revision_tags = {qn("w:" + name) for name in ("ins", "del", "moveFrom", "moveTo")}
            has_revisions = any(node.tag in revision_tags for node in doc.element.iter())
            observed_boxes = set()
            has_boxes = empty_boxes = False

            def visit():
                nonlocal visited
                visited += 1
                if visited > MAX_DOCX_BLOCKS:
                    raise ValueError("DOCX structural work limit exceeded")

            has_alternatives = prepare_textboxes(doc.element, visit)

            def note_id(node):
                raw = node.get(qn("w:id"))
                if not isinstance(raw, str) or not re.fullmatch(r"[+-]?[0-9]{1,10}", raw):
                    raise ValueError("DOCX note definition/reference is invalid")
                value = int(raw)
                if not -(2 ** 31) <= value < 2 ** 31:
                    raise ValueError("DOCX note definition/reference is invalid")
                return value

            def note_references(paragraph, locator):
                for ref in paragraph._p.iter(*note_tags):
                    visit()
                    # A box's own paragraph gets its own exact context below.
                    owner = next((a for a in ref.iterancestors() if a.tag == qn("w:p")), None)
                    if owner is not paragraph._p:
                        continue
                    run = ref.getparent()
                    parent = run.getparent()
                    if (run.tag != qn("w:r") or not (parent is paragraph._p or
                            (parent.tag == qn("w:hyperlink") and parent.getparent() is paragraph._p))):
                        raise ValueError("DOCX note reference context is unsupported")
                    references.append((note_tags[ref.tag], note_id(ref), locator))
                    observed_references.add(ref)

            def paragraph_boxes(paragraph, locator, table_depth):
                nonlocal has_boxes, empty_boxes
                bi = 0
                for box in paragraph._p.iter(BOX):
                    visit()
                    validate_box_context(box, paragraph._p)
                    observed_boxes.add(box)
                    has_boxes = True
                    bi += 1
                    prefix = f"{locator}, metin kutusu {bi}"
                    parent = SimpleNamespace(part=paragraph.part)
                    pi = ti = 0
                    before = len(passages)
                    for block in box:
                        visit()
                        if block.tag == qn("w:p"):
                            pi += 1
                            paragraph_text(Paragraph(block, parent), f"{prefix}, paragraf {pi}", table_depth)
                        elif block.tag == qn("w:tbl"):
                            ti += 1
                            table_text(Table(block, parent), f"{prefix}, tablo {ti}", table_depth + 1)
                        elif block.tag not in {qn("w:bookmarkStart"), qn("w:bookmarkEnd"), qn("w:proofErr")}:
                            raise ValueError("DOCX text-box content is unsupported")
                    empty_boxes |= len(passages) == before

            def paragraph_text(paragraph, locator, table_depth=0):
                note_references(paragraph, locator)
                add(locator, paragraph.text)
                paragraph_boxes(paragraph, locator, table_depth)

            def table_text(table, locator, depth):
                nonlocal has_tables
                has_tables = True
                if depth > MAX_DOCX_TABLE_DEPTH:
                    raise ValueError("DOCX table nesting limit exceeded")
                # Physical XML cells avoid expanding grid spans or repeating a
                # merged origin through row.cells. They are not visual columns.
                for ri, row in enumerate(table._tbl.tr_lst, 1):
                    visit()
                    cells = []
                    for node in row.tc_lst:
                        visit()
                        cells.append(_Cell(node, table))
                    row_locator = f"{locator}, satır {ri}"
                    if not any(cell.tables for cell in cells):
                        texts = []
                        box_paragraphs = []
                        for ci, cell in enumerate(cells, 1):
                            paragraphs = []
                            for pi, inner in enumerate(cell.iter_inner_content(), 1):
                                visit()
                                note_references(inner, f"{row_locator}, hücre {ci}, paragraf {pi}")
                                box_paragraphs.append((inner, f"{row_locator}, hücre {ci}, paragraf {pi}"))
                                paragraphs.append(inner.text)
                            texts.append("\n".join(paragraphs))
                        add(row_locator, " | ".join(texts))
                        for inner, context in box_paragraphs:
                            paragraph_boxes(inner, context, depth)
                        continue
                    # A nested row cannot be flattened without relocating its
                    # conditions. Retain each cell's paragraph/table sequence.
                    for ci, cell in enumerate(cells, 1):
                        cell_locator = f"{row_locator}, hücre {ci}"
                        pi = ti = 0
                        for inner in cell.iter_inner_content():
                            visit()
                            if isinstance(inner, Paragraph):
                                pi += 1
                                paragraph_text(inner, f"{cell_locator}, paragraf {pi}", depth)
                            else:
                                ti += 1
                                table_text(inner, f"{cell_locator}, tablo {ti}", depth + 1)

            # Paragraph/table locators retain independent numbering, including
            # blank paragraphs. Only their emitted body sequence changes.
            for block in doc.iter_inner_content():
                visit()
                if isinstance(block, Paragraph):
                    paragraph_index += 1
                    paragraph_text(block, f"Paragraf {paragraph_index}")
                else:
                    table_index += 1
                    table_text(block, f"Tablo {table_index}", 1)

            # Resolve only explicit package relationships. Accessing a missing
            # python-docx header can create a new definition; do not synthesize
            # source content or recursively follow inherited section wrappers.
            definitions = {}
            inspected_parts = set()
            has_stories = False
            variants = {"default": "varsayılan", "first": "ilk sayfa", "even": "çift sayfa"}
            story_types = (("header", "üstbilgi", RT.HEADER, HeaderPart),
                           ("footer", "altbilgi", RT.FOOTER, FooterPart))
            for si, section in enumerate(doc.sections, 1):
                visit()
                for kind, label, relation_type, part_type in story_types:
                    seen = set()
                    for ref in section._sectPr.findall(qn("w:" + kind + "Reference")):
                        visit()
                        variant = ref.get(qn("w:type"))
                        if variant not in variants or variant in seen:
                            raise ValueError("DOCX header/footer reference is invalid")
                        seen.add(variant)
                        relation = doc.part.rels.get(ref.get(qn("r:id")))
                        if (relation is None or relation.is_external
                                or relation.reltype != relation_type):
                            raise ValueError("DOCX header/footer reference is invalid")
                        part = relation.target_part
                        if (not isinstance(part, part_type)
                                or part.element.tag != qn("w:hdr" if kind == "header" else "w:ftr")):
                            raise ValueError("DOCX header/footer reference is invalid")
                        # An explicitly empty part overrides prior content too.
                        definitions[kind, variant] = (part, si)
                    for variant, variant_label in variants.items():
                        if (kind, variant) not in definitions:
                            continue
                        visit()
                        has_stories = True
                        part, source_section = definitions[kind, variant]
                        if part not in inspected_parts:
                            has_revisions |= any(node.tag in revision_tags for node in part.element.iter())
                            has_alternatives |= prepare_textboxes(part.element, visit)
                            source_roots.append(part.element)
                            inspected_parts.add(part)
                        prefix = (f"Bölüm {si}, {label} tanımı "
                                  f"({variant_label}; kaynak bölüm {source_section})")
                        # Definitions follow the complete body as a separate
                        # appendix. Their locators never assert page placement.
                        container = BlockItemContainer(part.element, SimpleNamespace(part=part))
                        pi = ti = 0
                        for block in container.iter_inner_content():
                            visit()
                            if isinstance(block, Paragraph):
                                pi += 1
                                paragraph_text(block, f"{prefix}, paragraf {pi}")
                            else:
                                ti += 1
                                table_text(block, f"{prefix}, tablo {ti}", 1)

            # python-docx's relationship map can collapse duplicate XML IDs.
            # Reject that ambiguity before selecting either note authority part.
            with zipfile.ZipFile(path) as archive:
                rel_path = str(doc.part.partname.rels_uri).lstrip("/")
                if rel_path in archive.namelist():
                    rel_root = fromstring(archive.read(rel_path), forbid_dtd=True)
                    rel_ids = set()
                    for rel in rel_root:
                        visit()
                        rid = rel.get("Id")
                        if not rid or rid in rel_ids:
                            raise ValueError("DOCX note relationship is invalid")
                        rel_ids.add(rid)

            notes = {}
            note_labels = {"footnote": "Dipnot", "endnote": "Sonnot"}
            for kind in note_labels:
                relation_type = getattr(RT, kind.upper() + "S")
                relations = [r for r in doc.part.rels.values() if r.reltype == relation_type]
                if len(relations) > 1 or any(r.is_external for r in relations):
                    raise ValueError("DOCX note relationship is invalid")
                if not relations:
                    continue
                visit()
                part = relations[0].target_part
                if part.content_type != getattr(CT, "WML_" + kind.upper() + "S"):
                    raise ValueError("DOCX note relationship is invalid")
                # Validate hostile XML before constructing typed Word blocks.
                fromstring(part.blob, forbid_dtd=True)
                root = parse_xml(part.blob)
                if root.tag != qn("w:" + kind + "s"):
                    raise ValueError("DOCX note definition/reference is invalid")
                has_revisions |= any(node.tag in revision_tags for node in root.iter())
                # A note cannot contain another note, including in an unselected
                # shape alternative. Validate before normalization.
                if any(True for _ in root.iter(*note_tags)):
                    raise ValueError("DOCX nested note reference is prohibited")
                has_alternatives |= prepare_textboxes(root, visit)
                identifiers = set()
                for node in root:
                    visit()
                    if node.tag != qn("w:" + kind):
                        raise ValueError("DOCX note definition/reference is invalid")
                    identifier = note_id(node)
                    note_type = node.get(qn("w:type"), "normal")
                    if identifier in identifiers or note_type not in {
                        "normal", "separator", "continuationSeparator", "continuationNotice"
                    }:
                        raise ValueError("DOCX note definition/reference is invalid")
                    identifiers.add(identifier)
                    if any(True for _ in node.iter(*note_tags)):
                        raise ValueError("DOCX nested note reference is prohibited")
                    if note_type == "normal":
                        notes[kind, identifier] = (node, part)
                        source_roots.append(node)

            contexts = {}
            for ordinal, (kind, identifier, locator) in enumerate(references, 1):
                if (kind, identifier) not in notes:
                    raise ValueError("DOCX note definition/reference is invalid")
                contexts.setdefault((kind, identifier), []).append((ordinal, locator))
            unreferenced = empty_notes = False
            for key, (node, part) in notes.items():
                kind, identifier = key
                occurrences = contexts.get(key, [(None, None)])
                unreferenced |= key not in contexts
                for ordinal, locator in occurrences:
                    visit()
                    origin = f"başvuru {ordinal} ({locator})" if ordinal is not None else "başvuru yok"
                    prefix = f"{note_labels[kind]} kimlik {identifier}, {origin}"
                    parent = SimpleNamespace(part=part)
                    pi = ti = 0
                    before_count = len(passages)
                    for block in node:
                        visit()
                        if block.tag == qn("w:p"):
                            pi += 1
                            paragraph_text(Paragraph(block, parent), f"{prefix}, paragraf {pi}")
                        elif block.tag == qn("w:tbl"):
                            ti += 1
                            table_text(Table(block, parent), f"{prefix}, tablo {ti}", 1)
                        else:
                            raise ValueError("DOCX note content is unsupported")
                    empty_notes |= len(passages) == before_count
            # Never assign an omitted wrapper's text or references to a guessed
            # host. Every selected box and source reference must be accounted for.
            for root in source_roots:
                for node in root.iter(BOX, *note_tags):
                    visit()
                    if node.tag == BOX and node not in observed_boxes:
                        raise ValueError("DOCX text-box context is unsupported")
                    if node.tag in note_tags and node not in observed_references:
                        raise ValueError("DOCX note reference context is unsupported")
            if has_boxes:
                warnings.append(
                    "Metin kutuları bağlayıcı paragraf ve XML kutu sırasıyla ayrı pasajlarda işlendi; "
                    "görsel okuma sırası, sayfa konumu, taşan/gizlenen metin ve alan değerleri doğrulanmadı."
                )
            if has_alternatives:
                warnings.append(
                    "Metin kutusunun eşdeğer XML alternatiflerinden biri işlendi; yinelenen kopya "
                    "ayrı kanıt sayılmadı. Bu işlem Word görünümünü doğrulamaz."
                )
            if empty_boxes:
                warnings.append("Boş metin kutusu var; bu kutudan metin pasajı oluşturulmadı.")
            if notes:
                warnings.append(
                    "Dipnot/sonnot paket kimliği ve XML başvuru konumu korunur; başvuru sırası gösterilen "
                    "not numarası değildir. Notlar ayrı ekte tanım sırasıyla, her başvuru bağlamında işlenir."
                )
            if unreferenced:
                warnings.append("Başvurusu bulunmayan not tanımları ayrı gösterildi; gövdeyle ilişkisi kurulmadı.")
            if empty_notes:
                warnings.append("Boş not tanımı var; bu tanımdan metin pasajı oluşturulmadı.")
            if has_stories:
                warnings.append(
                    "Üstbilgi/altbilgi tanımları gövdeden sonra bölüm ve kaynak bölüm bilgisiyle işlendi; "
                    "miras alınan tanım her ilgili bölümde ayrı gösterilir. Gerçek sayfa yerleşimi, "
                    "ilk/çift sayfa etkinliği ve alan değerleri doğrulanmadı."
                )
            if has_tables:
                warnings.append(
                    "Tablolar XML satır/hücre sırasıyla işlendi; hücre numarası fiziksel XML hücresidir. "
                    "Birleştirilmiş hücrelerin anlamı, görsel sütun hizası ve okuma düzeni doğrulanmadı."
                )
            if has_revisions:
                warnings.append("İzlenen değişiklikler var; kabul edilmiş sürüm doğrulanmadı.")
            warnings.append(
                "Belgedeki dipnot/sonnot sayfa yerleşimi, gösterilen numaralar ve çizim/görünüm tamlığı doğrulanmadı."
            )
        else:
            from openpyxl import load_workbook

            book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            count = 0
            for sheet in book:
                for row in sheet:
                    values = []
                    for cell in row:
                        if cell.value is not None:
                            count += 1
                            if count > 100000:
                                raise ValueError("Spreadsheet cell limit exceeded")
                            values.append(f"{cell.coordinate}: {cell.value}")
                            if cell.data_type == "f":
                                warnings.append(
                                    "Formüller çalıştırılmadı; değerler bağımsız doğrulama gerektirir."
                                )
                    if values:
                        add(f"{sheet.title}!{row[0].row}", " | ".join(values))
            book.close()
    elif suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(path)
        if len(presentation.slides) > 500:
            raise ValueError("Slide limit exceeded")
        page_count = len(presentation.slides)
        for si, slide in enumerate(presentation.slides, 1):
            for index, shape in enumerate(slide.shapes, 1):
                if shape.has_text_frame:
                    add(f"Slayt {si}, şekil {index}", shape.text_frame.text)
                if shape.has_table:
                    for ri, row in enumerate(shape.table.rows, 1):
                        add(f"Slayt {si}, tablo {index}, satır {ri}", " | ".join(c.text for c in row.cells))
        warnings.append("Slayt notları, gruplar, grafikler, görseller ve gömülü nesneler işlenmedi.")
    elif suffix in {".odt", ".ods", ".odp"}:
        from defusedxml.ElementTree import fromstring

        with zipfile.ZipFile(path) as archive:
            root = fromstring(archive.read("content.xml"))
        text_ns = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
        table_ns = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"

        def odf_text(element):
            pieces = [element.text or ""]
            for child in element:
                if child.tag == text_ns + "s":
                    count = int(child.attrib.get(text_ns + "c", "1"))
                    if not 1 <= count <= 1000:
                        raise ValueError("OpenDocument space expansion limit exceeded")
                    pieces.append(" " * count)
                elif child.tag == text_ns + "tab":
                    pieces.append("\t")
                elif child.tag == text_ns + "line-break":
                    pieces.append("\n")
                else:
                    pieces.append(odf_text(child))
                pieces.append(child.tail or "")
            return "".join(pieces)

        index = 0
        for element in root.iter():
            if element.tag in {text_ns + "p", text_ns + "h"}:
                index += 1
                add(f"İçerik paragrafı {index}", odf_text(element))
            if any(
                key in element.attrib
                for key in (
                    table_ns + "number-rows-repeated",
                    table_ns + "number-columns-repeated",
                    table_ns + "formula",
                )
            ):
                warnings.append("Tekrarlanan hücreler çoğaltılmadı; formüller çalıştırılmadı.")
        warnings.append(
            "OpenDocument: paragraf sırası korunur; hücre adresleri, düzen, nesneler ve imza doğrulanmadı."
        )
    elif suffix in {".html", ".htm"}:
        add("HTML metni", html_text(path.read_text(encoding="utf-8-sig", errors="strict")))
        warnings.append(
            "HTML kaynak metni: CSS görünürlüğü ve okuma düzeni doğrulanmadı; kaynaklar indirilmedi."
        )
    elif suffix == ".eml":
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        for header in ("Subject", "From", "To", "Date", "Message-ID"):
            if message[header]:
                add(f"E-posta {header}", str(message[header]))
        parts = list(message.walk())
        if len(parts) > 1000:
            raise ValueError("Email part limit exceeded")
        body = message.get_body(preferencelist=("plain", "html"))
        if body is not None:
            content = body.get_content()
            if isinstance(content, str):
                add(
                    "E-posta gövdesi",
                    html_text(content) if body.get_content_type() == "text/html" else content,
                )
        else:
            warnings.append("E-posta düz metin/HTML gövdesi bulunamadı; ileti gövdesi işlenmedi.")
        omitted = sum(
            1 for part in parts if part.get_content_disposition() == "attachment" or part.get_filename()
        )
        warnings.append(
            f"E-posta ekleri ayrı yüklenmelidir; {omitted} ek işlenmedi. Başlıkların özgünlüğü doğrulanmadı."
        )
    elif suffix == ".rtf":
        from striprtf.striprtf import rtf_to_text

        content = path.read_bytes()
        if not content.startswith(b"{\\rtf"):
            raise ValueError("Invalid RTF header")
        # RTF control sequences carry the codepage; raw bytes survive this reversible decode.
        add("RTF metni", rtf_to_text(content.decode("latin-1"), errors="strict"))
        warnings.append("RTF düzeni, alanlar ve gömülü nesneler doğrulanmadı; nesneler çalıştırılmadı.")
    elif suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path), on_demand=True)
        count = 0
        try:
            for sheet in book.sheets():
                count += sheet.nrows * sheet.ncols
                if count > 100000:
                    raise ValueError("Spreadsheet cell limit exceeded")
                for ri in range(sheet.nrows):
                    values = [
                        f"Sütun {ci + 1}: {sheet.cell_value(ri, ci)}"
                        for ci in range(sheet.ncols)
                        if sheet.cell_value(ri, ci) != ""
                    ]
                    add(f"{sheet.name}!Satır {ri + 1}", " | ".join(values))
        finally:
            book.release_resources()
        warnings.append(
            "XLS: kaydedilmiş hücre değerleri; formüller, makrolar, grafikler ve nesneler işlenmedi."
        )
    elif suffix == ".msg":
        import olefile

        with olefile.OleFileIO(str(path)) as message:
            streams = message.listdir()
            if len(streams) > 5000 or sum(message.get_size(s) for s in streams) > 100 * 1024 * 1024:
                raise ValueError("MSG stream limit exceeded")
            fields = {
                "0037": "Konu",
                "0C1A": "Gönderen adı",
                "0C1F": "Gönderen adresi",
                "0E04": "Alıcı",
                "1000": "Gövde",
                "007D": "İleti başlıkları",
            }
            for key, label in fields.items():
                stream = f"__substg1.0_{key}001F"
                if message.exists(stream):
                    add(
                        f"MSG {label}",
                        message.openstream(stream).read().decode("utf-16-le", errors="strict").rstrip("\0"),
                    )
            if not message.exists("__substg1.0_1000001F"):
                warnings.append("MSG Unicode düz metin gövdesi yok; HTML/RTF/ANSI gövde işlenmedi.")
            attachment_names = {s[0] for s in streams if s[0].startswith("__attach")}
            warnings.append(
                f"MSG: {len(attachment_names)} ek işlenmedi; ekleri ayrı yükleyin. Başlıklar doğrulanmadı."
            )
    elif suffix == ".doc":
        if not path.read_bytes().startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
            raise ValueError("Invalid legacy Word header")
        result = subprocess.run(
            ["antiword", "-m", "UTF-8.txt", str(path)], check=True, capture_output=True, timeout=20
        )
        add("Eski Word metni", result.stdout.decode("utf-8", errors="strict"))
        warnings.append("DOC: düz metin; düzen, değişiklikler, makrolar ve gömülü nesneler doğrulanmadı.")
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDF requires a decrypted working copy")
        if len(reader.pages) > 500:
            raise ValueError("PDF page limit exceeded")
        page_count = len(reader.pages)
        ocr_pages = 0
        for index, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if not text.strip():
                if ocr_pages < 3:
                    ocr_pages += 1
                    try:
                        import pytesseract
                        from PIL import Image

                        with tempfile.TemporaryDirectory(prefix="pdf-page-") as directory:
                            prefix = str(Path(directory) / "page")
                            subprocess.run(
                                [
                                    "pdftoppm",
                                    "-f",
                                    str(index),
                                    "-l",
                                    str(index),
                                    "-singlefile",
                                    "-scale-to",
                                    "2000",
                                    "-png",
                                    str(path),
                                    prefix,
                                ],
                                check=True,
                                capture_output=True,
                                timeout=6,
                            )
                            with Image.open(prefix + ".png") as picture:
                                text = pytesseract.image_to_string(picture, lang="tur+eng", timeout=6)
                        warnings.append(f"Sayfa {index}: OCR kullanıldı; görüntüyle doğrulayın.")
                    except (OSError, RuntimeError, subprocess.SubprocessError):
                        warnings.append(f"Sayfa {index}: yerel OCR kullanılamadı.")
                if not text.strip():
                    warnings.append(f"Sayfa {index}: içerik işlenemedi; bu çalışmaya dahil edilmedi.")
            add(f"Sayfa {index}", text)
        warnings.append("PDF okuma sırası ve tablo düzeni insan incelemesi gerektirir.")
    elif suffix in (".png", ".jpg", ".jpeg", ".tiff", ".tif", ".bmp", ".webp", ".gif"):
        import pytesseract
        from PIL import Image

        Image.MAX_IMAGE_PIXELS = 30_000_000
        with Image.open(path) as picture:
            if picture.width * picture.height > Image.MAX_IMAGE_PIXELS:
                raise ValueError("Image pixel limit exceeded")
            picture.verify()
        with Image.open(path) as picture:
            page_count = getattr(picture, "n_frames", 1)
            if page_count > 1:
                warnings.append(
                    f"Çok sayfalı görsel: yalnızca ilk kare işlendi; {page_count - 1} kare dahil edilmedi."
                )
            text = pytesseract.image_to_string(picture, lang="tur+eng", timeout=40)
        add("Görsel OCR", text)
        warnings.append("OCR çıktısı kaynak görselle doğrulanmalıdır.")
    else:
        raise ValueError("Format is not qualified; preserve original and convert locally")
    if sum(len(p["text"]) for p in passages) > 2_000_000:
        raise ValueError("Extracted text limit exceeded")
    result = {
        "passages": passages,
        "warnings": list(dict.fromkeys(warnings)),
        "page_count": page_count,
        "passage_count": len(passages),
    }
    if suffix == ".docx":
        result["extraction_profile"] = DOCX_BODY_PROFILE
    return result


if __name__ == "__main__":
    resource.setrlimit(resource.RLIMIT_CPU, (50, 50))
    resource.setrlimit(resource.RLIMIT_FSIZE, (10 * 1024 * 1024, 10 * 1024 * 1024))
    try:
        result = extract(Path(sys.argv[1]), sys.argv[2])
        print(json.dumps(result, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({"error": type(exc).__name__, "message": str(exc)[:200]}))
        sys.exit(1)
